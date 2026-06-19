#!/usr/bin/env python3
"""
Step 8: Extract tag markers into a Taggable column.

Reads sources/<character>_step7.xlsx. If a row contains [~5] in Command or
Command Full, the new Taggable column is set to TRUE and [~5] is stripped
from Command and Command Full. [Tag] is stripped from Move Name and
Move Name Full.

Output: sources/<character>_step8.xlsx

Usage:
    python3 step_8_extract_tags.py
"""

import glob
import os
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font


TAG_PATTERN = re.compile(r'\s*\[~5\]')
NAME_TAG_PATTERN = re.compile(r'\s*\[Tag\]')


def parse_sections(ws):
    """Parse a worksheet into sections."""
    sections = []
    row = 1
    while row <= ws.max_row:
        cell = ws.cell(row=row, column=1)
        if cell.value and cell.font.bold:
            next_cell = ws.cell(row=row + 1, column=1) if row + 1 <= ws.max_row else None
            if next_cell and next_cell.value == 'Command' and next_cell.font.bold:
                heading = cell.value
                row += 1

                header_cells = []
                for col in range(1, ws.max_column + 1):
                    val = ws.cell(row=row, column=col).value
                    if val:
                        header_cells.append(val)
                row += 1

                data_rows = []
                while row <= ws.max_row:
                    first_cell = ws.cell(row=row, column=1)
                    if first_cell.value is None and all(
                        ws.cell(row=row, column=c).value is None
                        for c in range(1, len(header_cells) + 1)
                    ):
                        row += 1
                        break
                    if first_cell.font.bold:
                        break

                    row_dict = {}
                    for col_idx, col_name in enumerate(header_cells):
                        val = ws.cell(row=row, column=col_idx + 1).value
                        row_dict[col_name] = val or ''
                    data_rows.append(row_dict)
                    row += 1

                sections.append((heading, data_rows))
            else:
                row += 1
        else:
            row += 1

    return sections


def process_sections(sections):
    """Extract tags from rows. Returns count of tagged rows."""
    tag_count = 0

    for _, rows in sections:
        for row in rows:
            cmd = str(row.get('Command', ''))
            cmd_full = str(row.get('Command Full', ''))

            has_tag = bool(TAG_PATTERN.search(cmd) or TAG_PATTERN.search(cmd_full))

            if has_tag:
                tag_count += 1
                row['Taggable'] = 'TRUE'
                row['Command'] = TAG_PATTERN.sub('', cmd).strip()
                row['Command Full'] = TAG_PATTERN.sub('', cmd_full).strip()
                alt_cmds = str(row.get('Alt Commands', ''))
                if alt_cmds:
                    row['Alt Commands'] = TAG_PATTERN.sub('', alt_cmds).strip()
                move_name = str(row.get('Move Name', ''))
                move_name_full = str(row.get('Move Name Full', ''))
                row['Move Name'] = NAME_TAG_PATTERN.sub('', move_name).strip()
                row['Move Name Full'] = NAME_TAG_PATTERN.sub('', move_name_full).strip()

    return tag_count


OUTPUT_COLUMNS = [
    'Command', 'Command Full', 'Alt Commands', 'Move Name', 'Move Name Full', 'To Stance',
    'Speed', 'Speed Full', 'Block Adv', 'Block Adv Full', 'Hit Adv', 'Hit Adv Full',
    'Counter Hit Adv', 'Counter Hit Adv Full', 'Damage', 'Damage Sum', 'Damage Full',
    'Hit Range', 'Hit Range Full', 'Throw Type', 'Throw Escape', 'Properties', 'Notes',
    'Taggable', 'UUID', 'ML UUID', 'Parent UUID', 'Unmatched',
]

HIDDEN_COLUMNS = {'Command', 'Move Name', 'Damage Sum', 'Hit Range',
                  'Block Adv', 'Hit Adv', 'Counter Hit Adv', 'Speed'}


def autofit_columns(ws, header_row_numbers):
    """Auto-fit column widths. Hidden columns get width 0."""
    for col in ws.columns:
        col_letter = col[0].column_letter
        header_value = None
        for cell in col:
            if cell.row in header_row_numbers and cell.value:
                header_value = cell.value
                break
        if header_value in HIDDEN_COLUMNS:
            ws.column_dimensions[col_letter].width = 0
            ws.column_dimensions[col_letter].hidden = True
            continue
        max_len = 0
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max_len + 2


def find_characters():
    """Find all characters that have step7 xlsx files."""
    files = glob.glob('sources/*_step7.xlsx')
    return sorted(
        os.path.basename(f).replace('_step7.xlsx', '')
        for f in files
        if not os.path.basename(f).startswith('~$')
    )


def main():
    characters = find_characters()
    if not characters:
        print("No step7 files found in sources/")
        return
    print(f"Found characters: {characters}")

    for char in characters:
        input_path = f"sources/{char}_step7.xlsx"
        output_path = f"sources/{char}_step8.xlsx"

        print(f"\n=== {char.capitalize()} ===")
        wb_in = load_workbook(input_path)
        ws_in = wb_in.active

        sections = parse_sections(ws_in)
        tag_count = process_sections(sections)

        print(f"  {tag_count} rows tagged")

        wb_out = Workbook()
        ws = wb_out.active
        ws.title = "Merged"
        row_num = 1
        header_row_numbers = set()

        for section_name, rows in sections:
            cell = ws.cell(row=row_num, column=1, value=section_name)
            cell.font = Font(bold=True)
            row_num += 1

            for col, h in enumerate(OUTPUT_COLUMNS, 1):
                cell = ws.cell(row=row_num, column=col, value=h)
                cell.font = Font(bold=True)
            header_row_numbers.add(row_num)
            row_num += 1

            for row_dict in rows:
                for col, col_name in enumerate(OUTPUT_COLUMNS, 1):
                    val = row_dict.get(col_name, '')
                    if val:
                        ws.cell(row=row_num, column=col, value=val)
                row_num += 1

            row_num += 1

        autofit_columns(ws, header_row_numbers)
        wb_out.save(output_path)
        print(f"  Written to {output_path}")

    print(f"\nDone. Processed {len(characters)} characters.")


if __name__ == '__main__':
    main()
