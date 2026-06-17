#!/usr/bin/env python3
"""
Validate that multi-hit rows can be correctly split by step 9.

Checks that the number of presses (space-separated tokens in Command Full)
matches the number of Block Adv Full values.

Usage:
    python3 tools/validate_split.py
"""

import re
from openpyxl import load_workbook

DIRECTIONS = {'f', 'b', 'd', 'u', 'df', 'db', 'uf', 'ub', 'N', 'F', 'B', 'D', 'U',
              'DF', 'DB', 'UF', 'UB', 'FC'}


def token_has_button(tok):
    return bool(re.search(r'[1-4]', tok))


def add_spaces(cmd):
    """Insert spaces between presses in a command string.

    Rules:
    - < always separates presses (delay)
    - ~ separates if followed by a token containing a button
    - , separates if prev token has a button (and doesn't end with bare direction)
      and next token has a button (and isn't direction+button after a bare direction ending)
    """
    trailing = ''
    m = re.search(r'(\s+-\s*(\[.*?\])?\s*)$', cmd)
    if m:
        trailing = m.group(0)
        cmd_clean = cmd[:m.start()]
    else:
        cmd_clean = cmd
    result = []
    i = 0
    while i < len(cmd_clean):
        ch = cmd_clean[i]
        if ch == '<':
            result.append(' ')
            result.append(ch)
        elif ch == '~':
            rest = cmd_clean[i+1:]
            next_sep = re.search(r'[,<~]', rest)
            next_part = rest[:next_sep.start()] if next_sep else rest
            if token_has_button(next_part):
                result.append(' ')
            result.append(ch)
        elif ch == ',':
            before = ''.join(result).split(' ')[-1] if result else ''
            rest = cmd_clean[i+1:]
            next_sep = re.search(r'[,<~]', rest)
            next_part = rest[:next_sep.start()] if next_sep else rest
            prev_has_btn = token_has_button(before)
            prev_ends_dir = False
            if '~' not in before:
                parts = before.split('<')
                last = parts[-1]
                if last in DIRECTIONS:
                    prev_ends_dir = True
            curr_is_dir_btn = bool(re.match(r'^[fbudFBUDN][fbudFBUDN]*\+[1-4]', next_part))
            curr_is_motion = '~' in next_part
            curr_has_btn = token_has_button(next_part)
            prev_is_pure_dir = not token_has_button(before)
            curr_is_pure_dir = not curr_has_btn
            if prev_is_pure_dir and curr_is_pure_dir:
                result.append(ch)
            elif prev_is_pure_dir and curr_is_dir_btn:
                result.append(ch)
            elif prev_is_pure_dir and curr_is_motion:
                result.append(ch)
            elif prev_is_pure_dir and curr_has_btn and not curr_is_dir_btn:
                result.append(' ')
                result.append(ch)
            elif prev_has_btn and prev_ends_dir:
                result.append(ch)
            elif prev_has_btn and not prev_ends_dir and curr_has_btn:
                result.append(' ')
                result.append(ch)
            elif prev_has_btn and curr_is_pure_dir:
                result.append(ch)
            else:
                result.append(ch)
        else:
            result.append(ch)
        i += 1
    final = ''.join(result).strip() + trailing
    final = re.sub(r'  +', ' ', final)
    return final


def count_presses(spaced_cmd):
    """Count presses from space-separated command. Ignores WS/SS/RDS/FC prefix and trailing ' -'."""
    clean = re.sub(r'\s+-\s*(\[.*?\])?\s*$', '', spaced_cmd)
    clean = re.sub(r'^(WS|SS|RDS|FC)\s+', '', clean)
    parts = clean.split(' ')
    return len(parts)


DONT_SPLIT = {
    ('jin', 'd+3+4'),
    ('ling', 'd+1'),
    ('ling', 'D+1'),
    ('ling', 'u+1+2'),
    ('ling', 'SS 4 [~B]'),
}
IGNORE = {('ling', 'FC,DF+2')}


def main():
    import glob
    import os

    files = glob.glob('sources/*_step8.xlsx')
    characters = sorted(
        os.path.basename(f).replace('_step8.xlsx', '')
        for f in files
        if not os.path.basename(f).startswith('~$')
    )

    ok_count = 0
    problems_less = []  # presses < block (multi-hit press)
    problems_more = []  # presses > block (todo cases)

    for char in characters:
        path = f'sources/{char}_step8.xlsx'
        wb = load_workbook(path)
        ws = wb.active
        row = 1
        while row <= ws.max_row:
            cell = ws.cell(row=row, column=1)
            if cell.value and cell.font.bold:
                next_cell = ws.cell(row=row+1, column=1) if row+1 <= ws.max_row else None
                if next_cell and next_cell.value and next_cell.font.bold:
                    row += 1
                    headers = []
                    for col in range(1, ws.max_column+1):
                        v = ws.cell(row=row, column=col).value
                        if v:
                            headers.append(v)
                    row += 1
                    while row <= ws.max_row:
                        first = ws.cell(row=row, column=1)
                        if first.value is None and all(ws.cell(row=row, column=c).value is None for c in range(1, len(headers)+1)):
                            row += 1
                            break
                        if first.font.bold:
                            break
                        rd = {}
                        for ci, cn in enumerate(headers):
                            rd[cn] = ws.cell(row=row, column=ci+1).value or ''

                        cmd = str(rd.get('Command Full', ''))
                        block = str(rd.get('Block Adv Full', '') or rd.get('Block Adv', ''))
                        block_parts = block.split()

                        if len(block_parts) > 1 and (char, cmd) not in DONT_SPLIT and (char, cmd) not in IGNORE:
                            new_cmd = add_spaces(cmd)
                            press_count = count_presses(new_cmd)
                            block_count = len(block_parts)

                            if press_count == block_count:
                                ok_count += 1
                            elif press_count < block_count:
                                problems_less.append((char, cmd, new_cmd, press_count, block_count, block))
                            else:
                                problems_more.append((char, cmd, new_cmd, press_count, block_count, block))

                        row += 1
                else:
                    row += 1
            else:
                row += 1

    total = ok_count + len(problems_less) + len(problems_more)
    print(f'OK: {ok_count}/{total}')
    print(f'Multi-hit press (presses < block): {len(problems_less)}')
    print(f'Todo cases (presses > block): {len(problems_more)}')
    print()

    if problems_less:
        print('MULTI-HIT PRESS (one press makes multiple hits):')
        for char, cmd, new_cmd, pc, bc, block in problems_less:
            print(f'  {char:6s} {cmd:35s} spaced: {new_cmd:40s} presses={pc} block={bc}')
        print()

    if problems_more:
        print('TODO CASES (presses > block):')
        for char, cmd, new_cmd, pc, bc, block in problems_more:
            print(f'  {char:6s} {cmd:35s} spaced: {new_cmd:40s} presses={pc} block={bc}')


if __name__ == '__main__':
    main()
