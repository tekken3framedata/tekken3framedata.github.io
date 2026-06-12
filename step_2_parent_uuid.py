#!/usr/bin/env python3
"""
Step 2: Normalize headers and add Parent UUID column.

Reads sources/<character>_step1.xlsx:
  - Unifies column headers across all sections within each sheet
  - Adds Parent UUID column based on continuation level (– markers)

Output: sources/<character>_step2.xlsx

Usage:
    python3 step_2_parent_uuid.py
"""

import glob
import os
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font


# Canonical column order per sheet type.
# All sections in a sheet will have these columns (even if empty for that section).
FD_COLUMNS = ['Command', 'Hit', 'Block Adv', 'Hit Adv', 'Counter Hit Adv', 'Notes', 'UUID', 'Parent UUID']
ML_COLUMNS = ['Command', 'Move Name', 'Stance', 'Damage', 'Hit Range', 'Throw Type', 'Throw Escape', 'Properties', 'Notes', 'UUID', 'Parent UUID']

HEADER_ALIASES = {
    'Throw Name': 'Move Name',
    'Type': 'Throw Type',
    'Escape': 'Throw Escape',
}


def get_continuation_level(command):
    """Determine the continuation level from leading en-dashes.

    Returns 0 for root moves, 1+ for continuations.
    '– X' = level 1, '–– X' = level 2, etc.
    """
    if not command:
        return 0
    match = re.match(r'^(–+)', command)
    if not match:
        return 0
    return len(match.group(1))


def parse_sections(ws):
    """Parse a step1 worksheet into sections.

    Returns list of (heading, header_row, data_rows) tuples.
    Each data_row is a dict mapping column name -> value.
    """
    sections = []
    row = 1
    while row <= ws.max_row:
        cell = ws.cell(row=row, column=1)
        if cell.value and cell.font.bold:
            heading = cell.value
            row += 1

            header_cells = []
            for col in range(1, ws.max_column + 1):
                val = ws.cell(row=row, column=col).value
                if val:
                    header_cells.append(val)
            row += 1

            data_rows = []
            while row <= ws.max_row:
                first_cell = ws.cell(row=row, column=1)
                if first_cell.value is None and all(
                    ws.cell(row=row, column=c).value is None
                    for c in range(1, len(header_cells) + 1)
                ):
                    row += 1
                    break
                if first_cell.font.bold:
                    break

                row_dict = {}
                for col_idx, col_name in enumerate(header_cells):
                    val = ws.cell(row=row, column=col_idx + 1).value
                    canonical = HEADER_ALIASES.get(col_name, col_name)
                    row_dict[canonical] = val or ''
                data_rows.append(row_dict)
                row += 1

            sections.append((heading, header_cells, data_rows))
        else:
            row += 1

    return sections


def add_parent_uuids(data_rows):
    """Add Parent UUID to each row based on continuation level."""
    uuid_by_level = {}
    for row in data_rows:
        cmd = row.get('Command', '')
        level = get_continuation_level(cmd)
        row_uuid = row.get('UUID', '')

        if level == 0:
            uuid_by_level = {0: row_uuid}
            row['Parent UUID'] = ''
        else:
            parent_level = level - 1
            row['Parent UUID'] = uuid_by_level.get(parent_level, '')
            uuid_by_level[level] = row_uuid

    return data_rows


def write_unified_sheet(wb, sheet_title, sections, columns):
    """Write sections with unified column headers to a new sheet."""
    ws = wb.create_sheet(sheet_title)
    row_num = 1

    for heading, _, data_rows in sections:
        # Section heading
        cell = ws.cell(row=row_num, column=1, value=heading)
        cell.font = Font(bold=True)
        row_num += 1

        # Unified header row
        for col, h in enumerate(columns, 1):
            cell = ws.cell(row=row_num, column=col, value=h)
            cell.font = Font(bold=True)
        row_num += 1

        # Data rows
        add_parent_uuids(data_rows)
        for row_dict in data_rows:
            for col, col_name in enumerate(columns, 1):
                val = row_dict.get(col_name, '')
                if val:
                    ws.cell(row=row_num, column=col, value=val)
            row_num += 1

        row_num += 1  # blank row

    return ws


def autofit_columns(ws):
    """Auto-fit all column widths based on cell content."""
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max_len + 2


def get_columns_for_sheet(sheet_title):
    """Return the canonical column list for a sheet type."""
    if sheet_title == 'Frame Data':
        return FD_COLUMNS
    return ML_COLUMNS


def find_characters():
    """Find all characters that have step1 xlsx files."""
    files = glob.glob('sources/*_step1.xlsx')
    return sorted(
        os.path.basename(f).replace('_step1.xlsx', '')
        for f in files
        if not os.path.basename(f).startswith('~$')
    )


def main():
    characters = find_characters()
    if not characters:
        print("No step1 files found in sources/")
        return
    print(f"Found characters: {characters}")

    for char in characters:
        input_path = f"sources/{char}_step1.xlsx"
        output_path = f"sources/{char}_step2.xlsx"

        print(f"\n=== {char.capitalize()} ===")
        wb_in = load_workbook(input_path)
        wb_out = Workbook()
        wb_out.remove(wb_out.active)

        for ws_in in wb_in.worksheets:
            print(f"  Processing sheet: {ws_in.title}")
            sections = parse_sections(ws_in)
            columns = get_columns_for_sheet(ws_in.title)
            ws_out = write_unified_sheet(wb_out, ws_in.title, sections, columns)
            autofit_columns(ws_out)

        wb_out.save(output_path)
        print(f"  Written to {output_path}")

    print(f"\nDone. Processed {len(characters)} characters.")


if __name__ == '__main__':
    main()
