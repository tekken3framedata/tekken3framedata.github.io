#!/usr/bin/env python3
"""
Step 3: Merge Frame Data and Movelist into a single table.

Reads sources/<character>_step2.xlsx and merges the two sheets by matching
commands within the same section. Match condition: exact command string match.

Output: sources/<character>_step3.xlsx (single sheet)

Usage:
    python3 step_3_merge.py
"""

import glob
import os
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font


MERGED_COLUMNS = [
    'Command', 'Move Name', 'Stance', 'Hit', 'Block Adv', 'Hit Adv', 'Counter Hit Adv',
    'Damage', 'Hit Range', 'Throw Type', 'Throw Escape', 'Properties', 'Notes',
    'UUID', 'ML UUID', 'Parent UUID', 'Unmatched',
]


def parse_sections(ws):
    """Parse a step2 worksheet into sections.

    Returns list of (heading, data_rows) tuples.
    Each data_row is a dict mapping column name -> value.
    """
    sections = []
    row = 1
    while row <= ws.max_row:
        cell = ws.cell(row=row, column=1)
        if cell.value and cell.font.bold:
            heading = cell.value
            row += 1

            # Read header row
            header_cells = []
            for col in range(1, ws.max_column + 1):
                val = ws.cell(row=row, column=col).value
                if val:
                    header_cells.append(val)
            row += 1

            # Read data rows
            data_rows = []
            while row <= ws.max_row:
                first_cell = ws.cell(row=row, column=1)
                if first_cell.value is None and all(
                    ws.cell(row=row, column=c).value is None
                    for c in range(1, len(header_cells) + 1)
                ):
                    row += 1
                    break
                if first_cell.font.bold:
                    break

                row_dict = {}
                for col_idx, col_name in enumerate(header_cells):
                    val = ws.cell(row=row, column=col_idx + 1).value
                    row_dict[col_name] = val or ''
                data_rows.append(row_dict)
                row += 1

            sections.append((heading, data_rows))
        else:
            row += 1

    return sections


def merge_section(fd_rows, ml_rows):
    """Merge FD and ML rows by exact command match.

    Returns list of merged row dicts and match statistics.
    """
    ml_by_cmd = {}
    for ml_row in ml_rows:
        cmd = ml_row.get('Command', '')
        if cmd not in ml_by_cmd:
            ml_by_cmd[cmd] = []
        ml_by_cmd[cmd].append(ml_row)

    merged = []
    matched_count = 0
    ml_used = set()

    for fd_row in fd_rows:
        cmd = fd_row.get('Command', '')
        row = dict(fd_row)

        candidates = ml_by_cmd.get(cmd, [])
        ml_match = None
        for i, ml_row in enumerate(candidates):
            key = (cmd, id(ml_row))
            if key not in ml_used:
                ml_match = ml_row
                ml_used.add(key)
                break

        if ml_match:
            matched_count += 1
            row['Move Name'] = ml_match.get('Move Name', '')
            row['Stance'] = ml_match.get('Stance', '')
            row['Damage'] = ml_match.get('Damage', '')
            row['Hit Range'] = ml_match.get('Hit Range', '')
            row['Throw Type'] = ml_match.get('Throw Type', '')
            row['Throw Escape'] = ml_match.get('Throw Escape', '')
            row['Properties'] = ml_match.get('Properties', '')
            if ml_match.get('Notes'):
                existing = row.get('Notes', '')
                ml_notes = ml_match['Notes']
                row['Notes'] = '; '.join(filter(None, [existing, ml_notes]))
            row['ML UUID'] = ml_match.get('UUID', '')
            row['Unmatched'] = ''
        else:
            row['Unmatched'] = 'FD only'

        merged.append(row)

    # ML rows with no FD match
    for ml_row in ml_rows:
        cmd = ml_row.get('Command', '')
        key = (cmd, id(ml_row))
        if key not in ml_used:
            row = {
                'Command': cmd,
                'Move Name': ml_row.get('Move Name', ''),
                'Stance': ml_row.get('Stance', ''),
                'Damage': ml_row.get('Damage', ''),
                'Hit Range': ml_row.get('Hit Range', ''),
                'Throw Type': ml_row.get('Throw Type', ''),
                'Throw Escape': ml_row.get('Throw Escape', ''),
                'Properties': ml_row.get('Properties', ''),
                'Notes': ml_row.get('Notes', ''),
                'ML UUID': ml_row.get('UUID', ''),
                'Parent UUID': ml_row.get('Parent UUID', ''),
                'Unmatched': 'ML only',
            }
            merged.append(row)

    return merged, matched_count, len(fd_rows), len(ml_rows)


def autofit_columns(ws):
    """Auto-fit all column widths based on cell content."""
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max_len + 2


def find_characters():
    """Find all characters that have step2 xlsx files."""
    files = glob.glob('sources/*_step2.xlsx')
    return sorted(
        os.path.basename(f).replace('_step2.xlsx', '')
        for f in files
        if not os.path.basename(f).startswith('~$')
    )


def main():
    characters = find_characters()
    if not characters:
        print("No step2 files found in sources/")
        return
    print(f"Found characters: {characters}")

    for char in characters:
        input_path = f"sources/{char}_step2.xlsx"
        output_path = f"sources/{char}_step3.xlsx"

        print(f"\n=== {char.capitalize()} ===")
        wb_in = load_workbook(input_path)

        fd_sections = parse_sections(wb_in['Frame Data'])
        ml_sections = parse_sections(wb_in['Movelist'])

        ml_by_name = {name: rows for name, rows in ml_sections}

        wb_out = Workbook()
        ws = wb_out.active
        ws.title = "Merged"
        row_num = 1

        all_sections = []
        fd_section_names = [name for name, _ in fd_sections]
        ml_only_sections = [name for name, _ in ml_sections if name not in fd_section_names]

        for section_name, fd_rows in fd_sections:
            ml_rows = ml_by_name.get(section_name, [])
            merged, matched, fd_count, ml_count = merge_section(fd_rows, ml_rows)

            unmatched_fd = sum(1 for r in merged if r.get('Unmatched') == 'FD only')
            unmatched_ml = sum(1 for r in merged if r.get('Unmatched') == 'ML only')
            print(f"  {section_name}: {matched}/{fd_count} FD matched, {unmatched_fd} FD only, {unmatched_ml} ML only")

            all_sections.append((section_name, merged))

        for section_name in ml_only_sections:
            ml_rows = ml_by_name[section_name]
            merged = []
            for ml_row in ml_rows:
                row = {
                    'Command': ml_row.get('Command', ''),
                    'Move Name': ml_row.get('Move Name', ''),
                    'Stance': ml_row.get('Stance', ''),
                    'Damage': ml_row.get('Damage', ''),
                    'Hit Range': ml_row.get('Hit Range', ''),
                    'Throw Type': ml_row.get('Throw Type', ''),
                    'Throw Escape': ml_row.get('Throw Escape', ''),
                    'Properties': ml_row.get('Properties', ''),
                    'Notes': ml_row.get('Notes', ''),
                    'ML UUID': ml_row.get('UUID', ''),
                    'Parent UUID': ml_row.get('Parent UUID', ''),
                    'Unmatched': 'ML only',
                }
                merged.append(row)
            print(f"  {section_name}: ML only ({len(ml_rows)} rows)")
            all_sections.append((section_name, merged))

        # Write output
        for section_name, merged in all_sections:
            cell = ws.cell(row=row_num, column=1, value=section_name)
            cell.font = Font(bold=True)
            row_num += 1

            for col, h in enumerate(MERGED_COLUMNS, 1):
                cell = ws.cell(row=row_num, column=col, value=h)
                cell.font = Font(bold=True)
            row_num += 1

            for row_dict in merged:
                for col, col_name in enumerate(MERGED_COLUMNS, 1):
                    val = row_dict.get(col_name, '')
                    if val:
                        ws.cell(row=row_num, column=col, value=val)
                row_num += 1

            row_num += 1  # blank row

        autofit_columns(ws)
        wb_out.save(output_path)
        print(f"  Written to {output_path}")

    print(f"\nDone. Processed {len(characters)} characters.")


if __name__ == '__main__':
    main()
