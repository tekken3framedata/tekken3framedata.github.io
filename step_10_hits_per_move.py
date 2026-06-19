#!/usr/bin/env python3
"""
Step 10: Add "Hits Per Move" column indicating how many hits each press produces.

Reads sources/<character>_step9.xlsx. Adds a "Hits Per Move" column where each
space-separated value corresponds to a press in Command Full. Default is "1" per
press. Multi-hit moves get their actual hit count.

Example: Command Full="1+4 ,2" with 1+4 being a 2-hit move → Hits Per Move="2 1"

Output: sources/<character>_step10.xlsx

Usage:
    python3 step_10_hits_per_move.py
"""

import glob
import os

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font


# Per-press hit counts by UUID. Each entry maps press index (0-based) to hit count.
# Unlisted presses default to 1.
HITS_PER_MOVE = {
    # --- Jin ---
    # b+1+3 (High and Mid Attack Reversal)
    'accd5398-5cb2-461a-897c-7c36276e215c': {0: 0},
    # 1 <2 ,4 ~1+4 ,2 (Double Punch – Knee – White Heron)
    '9688517c-2ed0-4fed-a34d-76e10b004eeb': {3: 2},
    # 1 <2 ,4 ~1+4 ,2 ,4 (Double Punch – Knee – White Heron – Crescent Kick)
    '99016883-22f7-46e2-a6b8-ce5a61a0d420': {3: 2},
    # 1 <2 ,4 ~1+4 ,2 ,d+4 (Double Punch – Knee – White Heron – Crescent Sweep)
    '7fc69d38-3ebd-4e97-9de3-fd5c4f4bea32': {3: 2},
    # 1+4 ,2 (White Heron)
    '7727ab03-4896-4e56-a0cc-ef259d5609a1': {0: 2},
    # 1+4 ,2 ,4 (White Heron – Crescent Kick)
    '4f2d088d-12f5-4b05-ba54-bb7c40204ad8': {0: 2},
    # 1+4 ,2 ,d+4 (White Heron – Crescent Sweep)
    '16c7aa83-e339-43c8-a2f8-aa7ab167725f': {0: 2},
    # d+3+4 (Can Can)
    'd5bd8f24-b46e-46c4-b044-d58438dfe3be': {0: 2},
    # WS+1 ,2 - (Twin Pistons -)
    '92abbd98-841a-40a6-9f2c-f893c21e340e': {2: 0},
    # WS+2 - (Uppercut -)
    'b5c3b733-3b05-4b65-8aa7-137cfd48bec5': {1: 0},
    # SS+2 - (Tooth Fairy -)
    'c6e4830e-a3cb-460d-9b03-0c4fe13b1861': {1: 0},
    # f,N,d,df+2 - (Wind Godfist -)
    '770ef824-eb8e-49ff-828c-c73385a322d8': {1: 0},
    # f~N,d~df+2 - (Electric Wind Godfist -)
    'ae1401a6-1714-44f2-8821-a0baa135dd79': {1: 0},
    # b,f+2 <1 <d+2 - (Laser Scraper -)
    '99c069f9-a4c3-4982-9d97-b5409dcd7cdd': {3: 0},
    # 4 ~3 (Demon Scissors)
    '198d48e3-fb8e-4387-b90f-d12b4dbd87ad': {0: 0},

    # --- Lee ---
    # b+3 ~3 (Feint Mist Wolf)
    'a1226d52-19ca-4e28-a360-74af905c7295': {0: 0},

    # --- Julia ---
    # 1 ~2 ~1 (G-Clef - Gut Punch – Skyscraper Cannon)
    '6275a354-d231-4442-8af8-0fc378c274f5': {2: 0},
    # f+1 ~2 (Palm Explosion)
    'a4c9455c-b3f7-445f-bb03-bc9714978245': {0: 0},
    # df+2 ,1 (Gut Punch – Skyscraper Cannon)
    'c672c34b-05e9-4636-968c-72c907a897d8': {1: 0},
    # 3+4 (Counter Clockwise Spin)
    'ad154462-6148-4f1a-8a02-ea05b3eeb7b2': {0: 0},
    # FC+1+3 (Low Parry)
    '378397f9-a5c6-48bc-b260-44dfff87353f': {0: 0},
    # 1+4 ,3 (Club Fist - Bow - Arrow)
    'ecf9ec9c-b22c-4f6a-9f5c-8f0b3d51b517': {0: 2},

    # --- Ling ---
    # f+1+2 (Cartwheel Dodge)
    '4579fe1a-4ded-4ab1-8378-2eec09ff10b8': {0: 0},
    # u+1+2 ,3+4 (Double Fan – Ginger Snap)
    'b929cc95-3142-4aba-8191-88bca5e84bf1': {0: 2, 1: 0},
    # 3+4 (Spinner Dodge)
    'd09dd708-e5fd-49ca-921c-de3f19c95bdb': {0: 0},
    # f+3+4 (Dive Roll)
    'fb27b361-92bd-42ac-894f-527ba7204cc5': {0: 0},
    # b+1+2 (Hypnotist Walk)
    'ce011a94-5f40-4d47-8059-3a12141afc50': {0: 0},
    # b+1+2 ,2 (Hypnotist Walk – Spin Sticker)
    '4fb4d6e5-1428-47b8-84f7-b084a82d6f57': {0: 0},
    # b+1+2,5 (Hypnotist Walk – Tag)
    'cde3cbb6-39e4-4bbd-8d1e-59d63f6766a9': {0: 0},
    # SS+4 [~B] (Twin Phoenix [Basic Stance])
    '575f73c7-dbfa-4d31-8a40-4db2663c3a4d': {0: 2},
    # AOP 4 ~3 [u_d] (Fire Cracker - [Roll up or down])
    'eba52a6a-4354-422e-bcd5-c23c11453761': {1: 0},
    # FC+db+3+4 (Crouching Rain Dance Starter)
    '9cd78cfe-f850-44ec-8cbc-c1744ce20f54': {0: 0},
    # d+1+2 (Phoenix Stance Starter)
    '3f6ba135-24e6-44ae-89c1-09dbdf045c6a': {0: 0},
    # b+3+4 (Rain Dance Starter)
    'def84375-5011-4b11-84a0-885a34b25d85': {0: 0},
    # u,ub (Evasive Backflip)
    '55f10e0e-b8fe-4478-8d4c-2fe839b19882': {0: 0},
    # 1+4 (High Parry)
    '0f3dc3b0-1a09-42da-8a39-ca793941d837': {0: 0},
    # d+1+4 (Low Parry)
    '66d0bc20-2aec-47e1-a748-5b027c60d425': {0: 0},
    # u+3+4 (Spin Cyclone)
    '72ceb8fa-8c13-414e-abde-e93840b782c9': {0: 0},
    # RDS d+1+2 (Art of Phoenix Stance)
    'e90f1b2c-7bb0-4304-acf8-33dbc4a1634a': {0: 0},
    # RDS 3+4 (Spinner Dodge)
    '83d45d20-abe6-4ff5-bd03-2cf9f2d0b03d': {0: 0},
    # RDS f+3+4 (California Roll)
    '228070f6-c47f-4f7b-a6b0-76d8abe818f3': {0: 0},
    # RDS f+3+4 ~3+4 (California Roll – Reverse Kangaroo Kick)
    'c6902a2f-fe58-4489-baf6-20e02a7ecb01': {0: 0},
    # AOP d (Butterfly)
    '8b15273f-5f2b-4cdd-b372-a139740596d7': {0: 0},
    # AOP f+3+4 (Roll Out)
    '987a1676-d0c0-438a-8e69-8ce9ec35af7f': {0: 0},
    # 1+3+4 (False Salute Taunt)
    '4f00b275-65c1-4e31-9c5b-02f3961f09dc': {0: 0},
    # 2+3+4 (Greetings Taunt) — 1 hit, 0 damage (default 1 is correct)
    # RDS 1+3+4 (False Salute Taunt)
    '07193dd0-f923-4b9e-bcdf-cfe43c6d0073': {0: 0},
    # b+1+2 <1+2 (Phoenix Strike)
    '25f9b418-8f7d-4dca-b2a8-04c676fbd4aa': {0: 0},
    # b+1+2 <1+2,b,b (Phoenix Strike – Cancel)
    '26909f52-7512-44a5-9bc0-0b5669dd6609': {1: 0},
    # d+1 (Flapping Wings)
    'a45a8361-3b4a-4e67-bc45-da8b314e5043': {0: 3},
    # D+1 (Alternate Flapping Wings)
    'eef6efe3-7a0b-4e48-8b25-1fdaf76cd199': {0: 3},
    # u+1+2 (Double Fan)
    '085e1996-b16b-4a63-a3d0-ae8152d065ea': {0: 2},
    # u+1+2 ,2 (Double Fan – Hydrangea)
    'aadbdd07-dac3-4f6e-9230-207fed9671ef': {0: 2},
    # u+1+2 ,2 ,1 (Double Fan – Fortune Cookie)
    '1fe5e046-832f-40fb-bc3f-b47d5d1f9f82': {0: 2},
}


def parse_presses(cmd_full):
    """Split Command Full into individual presses."""
    if not cmd_full or str(cmd_full).strip() == '':
        return []
    return str(cmd_full).split(' ')


def compute_hits_per_move(cmd_full, uuid=None):
    """Compute the hits-per-move string for a given Command Full."""
    presses = parse_presses(cmd_full)
    if not presses:
        return ''

    overrides = HITS_PER_MOVE.get(uuid, {}) if uuid else {}
    hits = []
    for i, press in enumerate(presses):
        if i in overrides:
            hits.append(str(overrides[i]))
        elif press.startswith('[') and press.endswith(']'):
            hits.append('0')
        else:
            hits.append('1')

    return ' '.join(hits)


def parse_sections(ws):
    """Parse a worksheet into sections."""
    sections = []
    row = 1
    while row <= ws.max_row:
        cell = ws.cell(row=row, column=1)
        if cell.value and cell.font.bold:
            next_cell = ws.cell(row=row + 1, column=1) if row + 1 <= ws.max_row else None
            if next_cell and next_cell.value == 'From Stance' and next_cell.font.bold:
                heading = cell.value
                row += 1

                header_cells = []
                for col in range(1, ws.max_column + 1):
                    val = ws.cell(row=row, column=col).value
                    if val:
                        header_cells.append(val)
                row += 1

                data_rows = []
                while row <= ws.max_row:
                    first_cell = ws.cell(row=row, column=1)
                    if first_cell.value is None and all(
                        ws.cell(row=row, column=c).value is None
                        for c in range(1, len(header_cells) + 1)
                    ):
                        row += 1
                        break
                    if first_cell.font.bold:
                        break

                    row_dict = {}
                    for col_idx, col_name in enumerate(header_cells):
                        val = ws.cell(row=row, column=col_idx + 1).value
                        row_dict[col_name] = val or ''
                    data_rows.append(row_dict)
                    row += 1

                sections.append((heading, header_cells, data_rows))
            else:
                row += 1
        else:
            row += 1

    return sections


def process_sections(sections):
    """Add Hits Per Move column to all rows."""
    for _, _, rows in sections:
        for row in rows:
            cmd_full = str(row.get('Command Full', '') or '')
            uuid = str(row.get('UUID', '') or '')
            row['Hits Per Move'] = compute_hits_per_move(cmd_full, uuid)


def autofit_columns(ws, header_row_numbers, hidden_columns):
    """Auto-fit column widths. Hidden columns get width 0."""
    for col in ws.columns:
        col_letter = col[0].column_letter
        header_value = None
        for cell in col:
            if cell.row in header_row_numbers and cell.value:
                header_value = cell.value
                break
        if header_value in hidden_columns:
            ws.column_dimensions[col_letter].width = 0
            ws.column_dimensions[col_letter].hidden = True
            continue
        max_len = 0
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max_len + 2


OUTPUT_COLUMNS = [
    'From Stance',
    'Command', 'Command Full', 'Alt Commands', 'Move Name', 'Move Name Full', 'To Stance',
    'Hits Per Move',
    'Speed', 'Speed Full', 'Block Adv', 'Block Adv Full', 'Hit Adv', 'Hit Adv Full',
    'Counter Hit Adv', 'Counter Hit Adv Full', 'Damage', 'Damage Sum', 'Damage Full',
    'Hit Range', 'Hit Range Full', 'Throw Type', 'Throw Escape', 'Properties', 'Notes',
    'Taggable', 'UUID', 'ML UUID', 'Parent UUID', 'Unmatched',
]

HIDDEN_COLUMNS = {'Command', 'Move Name', 'Damage Sum', 'Hit Range',
                  'Block Adv', 'Hit Adv', 'Counter Hit Adv', 'Speed'}


def find_characters():
    """Find all characters that have step9 xlsx files."""
    files = glob.glob('sources/*_step9.xlsx')
    return sorted(
        os.path.basename(f).replace('_step9.xlsx', '')
        for f in files
        if not os.path.basename(f).startswith('~$')
    )


def main():
    characters = find_characters()
    if not characters:
        print("No step9 files found in sources/")
        return
    print(f"Found characters: {characters}")

    for char in characters:
        input_path = f"sources/{char}_step9.xlsx"
        output_path = f"sources/{char}_step10.xlsx"

        print(f"\n=== {char.capitalize()} ===")
        wb_in = load_workbook(input_path)
        ws_in = wb_in.active

        sections = parse_sections(ws_in)
        process_sections(sections)

        multi_hit_count = sum(
            1 for _, _, rows in sections
            for row in rows
            if any(int(x) > 1 for x in row.get('Hits Per Move', '').split() if x.isdigit())
        )
        print(f"  {multi_hit_count} rows with multi-hit moves")

        wb_out = Workbook()
        ws = wb_out.active
        ws.title = "Merged"
        row_num = 1
        header_row_numbers = set()

        for section_name, header_cells, rows in sections:
            cell = ws.cell(row=row_num, column=1, value=section_name)
            cell.font = Font(bold=True)
            row_num += 1

            for col, h in enumerate(OUTPUT_COLUMNS, 1):
                cell = ws.cell(row=row_num, column=col, value=h)
                cell.font = Font(bold=True)
            header_row_numbers.add(row_num)
            row_num += 1

            for row_dict in rows:
                for col, col_name in enumerate(OUTPUT_COLUMNS, 1):
                    val = row_dict.get(col_name, '')
                    if val:
                        ws.cell(row=row_num, column=col, value=val)
                row_num += 1

            row_num += 1

        autofit_columns(ws, header_row_numbers, HIDDEN_COLUMNS)
        wb_out.save(output_path)
        print(f"  Written to {output_path}")

    print(f"\nDone. Processed {len(characters)} characters.")


if __name__ == '__main__':
    main()
