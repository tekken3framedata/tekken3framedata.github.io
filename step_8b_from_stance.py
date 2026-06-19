#!/usr/bin/env python3
"""
Step 8b: Add From Stance column based on section name and command prefix.

Reads sources/<character>_step8.xlsx. Adds a From Stance column derived from:
  1. Stance prefix in Command Full (e.g. "BT 1 ,2" → From Stance="BT", Command Full="1 ,2")
     Recognized prefixes: AOP, BT, FCD, LFS, PLD, RDS, RFS
  2. Section heading (for moves without a command prefix):
     - Rain Dance Art → RDS
     - Art Of Phoenix → AOP
     - Devil Jin Possession Arts → DJP

Output: sources/<character>_step8b.xlsx

Usage:
    python3 step_8b_from_stance.py
"""

import glob
import os

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font


STANCE_PREFIXES = ('AOP', 'BT', 'FCD', 'LFS', 'PLD', 'RDS', 'RFS')

SECTION_TO_STANCE = {
    'Rain Dance Art': 'RDS',
    'Art Of Phoenix': 'AOP',
    'Art of Phoenix Illusion': 'AOP',
    'Devil Jin Possession Arts': 'DJP',
}


def parse_sections(ws):
    sections = []
    row = 1
    while row <= ws.max_row:
        cell = ws.cell(row=row, column=1)
        if cell.value and cell.font.bold:
            next_cell = ws.cell(row=row + 1, column=1) if row + 1 <= ws.max_row else None
            if next_cell and next_cell.value == 'Command' and next_cell.font.bold:
                heading = cell.value
                row += 1

                header_cells = []
                for col in range(1, ws.max_column + 1):
                    val = ws.cell(row=row, column=col).value
                    if val:
                        header_cells.append(val)
                row += 1

                data_rows = []
                while row <= ws.max_row:
                    first_cell = ws.cell(row=row, column=1)
                    if first_cell.value is None and all(
                        ws.cell(row=row, column=c).value is None
                        for c in range(1, len(header_cells) + 1)
                    ):
                        row += 1
                        break
                    if first_cell.font.bold:
                        break

                    row_dict = {}
                    for col_idx, col_name in enumerate(header_cells):
                        val = ws.cell(row=row, column=col_idx + 1).value
                        row_dict[col_name] = val or ''
                    data_rows.append(row_dict)
                    row += 1

                sections.append((heading, data_rows))
            else:
                row += 1
        else:
            row += 1

    return sections


def extract_stance_prefix(cmd):
    """Extract stance prefix from command. Returns (stance, remaining_cmd)."""
    for stance in STANCE_PREFIXES:
        if cmd.startswith(stance + ' '):
            return stance, cmd[len(stance) + 1:]
    return '', cmd


def process_sections(sections):
    count = 0
    prefix_count = 0
    for i, (section_name, rows) in enumerate(sections):
        stance = SECTION_TO_STANCE.get(section_name, '')
        filtered = [
            row for row in rows
            if not str(row.get('Command Full', '') or '').startswith('Moves starting from')
        ]
        sections[i] = (section_name, filtered)
        for row in filtered:
            cmd_full = str(row.get('Command Full', '') or '')
            prefix, remaining = extract_stance_prefix(cmd_full)
            if prefix:
                row['From Stance'] = prefix
                row['Command Full'] = remaining
                cmd = str(row.get('Command', '') or '')
                cmd_prefix, cmd_remaining = extract_stance_prefix(cmd)
                if cmd_prefix:
                    row['Command'] = cmd_remaining
                prefix_count += 1
            elif not row.get('From Stance'):
                row['From Stance'] = stance
                if stance:
                    count += 1
    return count, prefix_count


OUTPUT_COLUMNS = [
    'From Stance',
    'Command', 'Command Full', 'Alt Commands', 'Move Name', 'Move Name Full', 'To Stance',
    'Speed', 'Speed Full', 'Block Adv', 'Block Adv Full', 'Hit Adv', 'Hit Adv Full',
    'Counter Hit Adv', 'Counter Hit Adv Full', 'Damage', 'Damage Sum', 'Damage Full',
    'Hit Range', 'Hit Range Full', 'Throw Type', 'Throw Escape', 'Properties', 'Notes',
    'Taggable', 'UUID', 'ML UUID', 'Parent UUID', 'Unmatched',
]

HIDDEN_COLUMNS = {'Command', 'Move Name', 'Damage Sum', 'Hit Range',
                  'Block Adv', 'Hit Adv', 'Counter Hit Adv', 'Speed'}


def autofit_columns(ws, header_row_numbers):
    for col in ws.columns:
        col_letter = col[0].column_letter
        header_value = None
        for cell in col:
            if cell.row in header_row_numbers and cell.value:
                header_value = cell.value
                break
        if header_value in HIDDEN_COLUMNS:
            ws.column_dimensions[col_letter].width = 0
            ws.column_dimensions[col_letter].hidden = True
            continue
        max_len = 0
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max_len + 2


def find_characters():
    files = glob.glob('sources/*_step8.xlsx')
    return sorted(
        os.path.basename(f).replace('_step8.xlsx', '')
        for f in files
        if not os.path.basename(f).startswith('~$')
    )


def main():
    characters = find_characters()
    if not characters:
        print("No step8 files found in sources/")
        return
    print(f"Found characters: {characters}")

    for char in characters:
        input_path = f"sources/{char}_step8.xlsx"
        output_path = f"sources/{char}_step8b.xlsx"

        print(f"\n=== {char.capitalize()} ===")
        wb_in = load_workbook(input_path)
        ws_in = wb_in.active

        sections = parse_sections(ws_in)
        count, prefix_count = process_sections(sections)

        print(f"  {count} rows from section, {prefix_count} rows from command prefix")

        wb_out = Workbook()
        ws = wb_out.active
        ws.title = "Merged"
        row_num = 1
        header_row_numbers = set()

        for section_name, rows in sections:
            cell = ws.cell(row=row_num, column=1, value=section_name)
            cell.font = Font(bold=True)
            row_num += 1

            for col, h in enumerate(OUTPUT_COLUMNS, 1):
                cell = ws.cell(row=row_num, column=col, value=h)
                cell.font = Font(bold=True)
            header_row_numbers.add(row_num)
            row_num += 1

            for row_dict in rows:
                for col, col_name in enumerate(OUTPUT_COLUMNS, 1):
                    val = row_dict.get(col_name, '')
                    if val:
                        ws.cell(row=row_num, column=col, value=val)
                row_num += 1

            row_num += 1

        autofit_columns(ws, header_row_numbers)
        wb_out.save(output_path)
        print(f"  Written to {output_path}")

    print(f"\nDone. Processed {len(characters)} characters.")


if __name__ == '__main__':
    main()
