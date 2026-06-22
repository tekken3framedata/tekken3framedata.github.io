#!/usr/bin/env python3
"""
Step 6b: Process bracket notations in Command column.

Reads sources/<character>_step6.xlsx. Extracts bracket notations and creates
new columns based on rules:
  - [~5]: Taggable=TRUE, strip from Command. Strip [Tag] from Move Name.

Output: sources/<character>_step6b.xlsx

Usage:
    python3 step_6b_bracket_notations.py
"""

import glob
import os
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font


TAG_PATTERN = re.compile(r'\s*\[~5\]')
NAME_TAG_PATTERN = re.compile(r'\s*\[Tag\]')

STANCE_RECOVERY = {
    'anna': [
        {
            'uuid': '68a6b41e-e87d-46fd-a7e1-f471e9ceb5c5',
            'strip_cmd': re.compile(r'\s*\[~U_~D\]'),
            'strip_name': re.compile(r'\s*\[SS\]'),
            'to_stance': '(~U)SS; (~D)SS',
        },
        {
            'uuid': '02266c87-12dd-41d3-879d-663dee145db8',
            'strip_cmd': re.compile(r'\s*\[~U_~D\]'),
            'strip_name': re.compile(r'\s*\[SS\]'),
            'to_stance': '(~U)SS; (~D)SS',
        },
        {
            'uuid': '37690d6b-180b-44cf-ac43-3c4eadd6b614',
            'strip_cmd': re.compile(r'\s*\[~U_~D\]'),
            'strip_name': re.compile(r'\s*\[SS\]'),
            'to_stance': '(~U)SS; (~D)SS',
        },
        {
            'uuid': 'd942c8fb-e8e1-4ae4-bdb6-38edab1b7d11',
            'strip_cmd': re.compile(r'\s*\[~U_~D\]'),
            'strip_name': re.compile(r'\s*\[SS\]'),
            'to_stance': '(~U)SS; (~D)SS',
        },
        {
            'uuid': '5a1a28e0-cd2d-4a6f-a14c-8a8ce7088ad1',
            'strip_cmd': re.compile(r'\s*\[~U_~D\]'),
            'strip_name': re.compile(r'\s*\[SS\]'),
            'to_stance': '(~U)SS; (~D)SS',
        },
        {
            'uuid': 'ab9308ff-e068-4c15-b29c-1171421322ea',
            'strip_cmd': re.compile(r'\s*\[~U_~D\]'),
            'strip_name': re.compile(r'\s*\[SS\]'),
            'to_stance': '(~U)SS; (~D)SS',
        },
        {
            'uuid': 'bbc2cf9a-42bd-4b05-93d8-4e31d4c0074c',
            'strip_cmd': re.compile(r'\s*\[~U_~D\]'),
            'strip_name': re.compile(r'\s*\[SS\]'),
            'to_stance': '(~U)SS; (~D)SS',
        },
    ],
    'baek': [
        {
            'uuid': '39d5078c-5045-43d1-984a-2fc1bc5573d3',
            'strip_cmd': re.compile(r'\s*\[~f_~b\]'),
            'strip_name': re.compile(r'\s*\[FLA\]'),
            'to_stance': '(~f)FLA; (~b)FLA',
        },
        {
            'uuid': 'f41fef9f-4c9e-473d-8021-337c2ce01a12',
            'strip_cmd': re.compile(r'\s*\[~f_~b\]'),
            'strip_name': re.compile(r'\s*\[FLA\]'),
            'to_stance': '(~f)FLA; (~b)FLA',
        },
        {
            'uuid': '65ddbdac-762e-457d-a52c-3bb4f27d6cfa',
            'strip_cmd': re.compile(r'\s*\[~f_~b\]'),
            'strip_name': re.compile(r'\s*\[FLA\]'),
            'to_stance': '(~f)FLA; (~b)FLA',
        },
        {
            'uuid': '2c2b3334-99a7-475e-b0b9-bebffc48cc17',
            'strip_cmd': re.compile(r'\s*\[~f_~b\]'),
            'strip_name': re.compile(r'\s*\[FLA\]'),
            'to_stance': '(~f)FLA; (~b)FLA',
        },
        {
            'uuid': 'd02dc538-7fa1-47ba-a87f-b003c3f9f8a2',
            'strip_cmd': re.compile(r'\s*\[~f_~b\]'),
            'strip_name': re.compile(r'\s*\[FLA\]'),
            'to_stance': '(~f)FLA; (~b)FLA',
        },
        {
            'uuid': '42ec062b-f6bf-461e-8847-fc894c0b3ed5',
            'strip_cmd': re.compile(r'\s*\[f_b\]'),
            'strip_name': re.compile(r'\s*\[FLA\]'),
            'to_stance': '(~F)FLA; (~B)FLA',
        },
        {
            'uuid': 'e917fbb5-b0df-4f3c-84c5-dbfd0a53d67d',
            'strip_cmd': re.compile(r'\s*\[f_b\]'),
            'strip_name': re.compile(r'\s*\[FLA\]'),
            'to_stance': '(~F)FLA; (~B)FLA',
        },
        {
            'uuid': 'ed571bf6-bc3b-4c92-b25b-29fc2afdc4b9',
            'strip_cmd': re.compile(r'\s*\[f_b\]'),
            'strip_name': re.compile(r'\s*\[FLA\]'),
            'to_stance': '(~F)FLA; (~B)FLA',
        },
        {
            'uuid': '03fbf9ce-64e4-44bb-bd8f-407890ad6887',
            'strip_cmd': re.compile(r'\s*\[f_b\]'),
            'strip_name': re.compile(r'\s*\[FLA\]'),
            'to_stance': '(~F)FLA; (~B)FLA',
        },
        {
            'uuid': 'aaf7ad1e-8140-4ab5-a9c8-b114edebebf1',
            'strip_cmd': re.compile(r'\s*\[f_b\]'),
            'strip_name': re.compile(r'\s*\[FLA\]'),
            'to_stance': '(~F)FLA; (~B)FLA',
        },
        {
            'uuid': '3a65dbc1-ca16-4e1c-914b-1a8529f4ba11',
            'strip_cmd': re.compile(r'\s*\[f_b\]'),
            'strip_name': re.compile(r'\s*\[FLA\]'),
            'to_stance': '(~F)FLA; (~B)FLA',
        },
        {
            'uuid': '40aa667d-0dc7-43f8-b2a3-7787fbb240d5',
            'strip_cmd': re.compile(r'\s*\[f_b\]'),
            'strip_name': re.compile(r'\s*\[FLA\]'),
            'to_stance': '(~F)FLA; (~B)FLA',
        },
        {
            'uuid': '5abeb8e1-894e-4fbc-8e4c-c044f37eb3d8',
            'strip_cmd': re.compile(r'\s*\[~B\]'),
            'strip_name': re.compile(r'\s*\[FLA\]'),
            'to_stance': '(~B)FLA',
        },
    ],
    'eddy': [
        {
            'uuid': 'b181e799-480d-4934-bffb-e68585285059',
            'strip_cmd': re.compile(r'\s*\[~B\]\s*\[~D\]'),
            'strip_name': re.compile(r'\s*\[HSP\s*\]\s*\[RLX\]'),
            'to_stance': '(~B)HSP; (~D)RLX',
        },
        {
            'uuid': '3398a120-61a1-4fc2-b813-34ff6ec1384c',
            'strip_cmd': re.compile(r'\s*\[~B\]'),
            'strip_name': re.compile(r'\s*\[HSP\]'),
            'to_stance': '(~B)HSP',
        },
        {
            'uuid': '6352e0bb-fc5e-48ad-8941-fc1b289167aa',
            'strip_cmd': re.compile(r'\s*\[~B\]'),
            'strip_name': re.compile(r'\s*\[HSP\]'),
            'to_stance': '(~B)HSP',
        },
        {
            'uuid': '4d3ee92e-c72e-4c16-a26d-664ba55f5469',
            'strip_cmd': re.compile(r'\s*\[~B\]'),
            'strip_name': re.compile(r'\s*\[HSP\]'),
            'to_stance': '(~B)HSP',
        },
        {
            'uuid': 'bec58a9f-3228-4419-babf-28b3cc3edb64',
            'strip_cmd': re.compile(r'\s*\[~B\]'),
            'strip_name': re.compile(r'\s*\[HSP\]'),
            'to_stance': '(~B)HSP',
        },
        {
            'uuid': '2fdf90ca-3648-460d-9ea8-a16f3d5f25e1',
            'strip_cmd': re.compile(r'\s*\[~B\]\s*\[~D\]'),
            'strip_name': re.compile(r'\s*\[HSP\]\s*\[RLX\]'),
            'to_stance': '(~B)HSP; (~D)RLX',
        },
        {
            'uuid': '8304ed6b-41b5-404d-a577-c0dfb41da1c8',
            'strip_cmd': re.compile(r'\s*\[~B\]'),
            'strip_name': re.compile(r'\s*\[HSP\]'),
            'to_stance': '(~B)HSP',
        },
        {
            'uuid': '9d075145-17ed-4f3b-a39c-c29cb0995dd2',
            'strip_cmd': re.compile(r'\s*\[~B\]'),
            'strip_name': re.compile(r'\s*\[HSP\]'),
            'to_stance': '(~B)HSP',
        },
        {
            'uuid': '14b1b2bb-a746-4688-af50-32c457a27e51',
            'strip_cmd': re.compile(r'\s*\[~B\]'),
            'strip_name': re.compile(r'\s*\[HSP\]'),
            'to_stance': '(~B)HSP',
        },
        {
            'uuid': 'c26557d6-15d4-466d-8bd2-40b947e2fa28',
            'strip_cmd': re.compile(r'\s*\[~B\]'),
            'strip_name': re.compile(r'\s*\[HSP\]'),
            'to_stance': '(~B)HSP',
        },
        {
            'uuid': 'e0c3444e-aae6-4eb9-9865-41b5d617f5e5',
            'strip_cmd': re.compile(r'\s*\[~B\]'),
            'strip_name': re.compile(r'\s*\[HSP\]'),
            'to_stance': '(~B)HSP',
        },
        {
            'uuid': 'c89588be-ef77-49f2-aa4b-f50702afd7b5',
            'strip_cmd': re.compile(r'\s*\[~B\]'),
            'strip_name': re.compile(r'\s*\[HSP\]'),
            'to_stance': '(~B)HSP',
        },
        {
            'uuid': '19930be2-c919-4d95-9092-ae4efe759060',
            'strip_cmd': re.compile(r'\s*\[~B\]'),
            'strip_name': re.compile(r'\s*\[HSP\]'),
            'to_stance': '(~B)HSP',
        },
        {
            'uuid': '635b15ef-4bc4-43e3-9185-c134f5ba4677',
            'strip_cmd': re.compile(r'\s*\[~D\]'),
            'strip_name': re.compile(r'\s*\[RLX\]'),
            'to_stance': '(~D)RLX',
        },
        {
            'uuid': '0f5c0eb7-d105-4c6d-9baa-b93ca74ce569',
            'strip_cmd': re.compile(r'\s*\[~B\]'),
            'strip_name': re.compile(r'\s*\[HSP\]'),
            'to_stance': '(~B)HSP',
        },
        {
            'uuid': '1f33d2bc-9841-4057-9510-bae5bf5e46f4',
            'strip_cmd': re.compile(r'\s*\[~B\]'),
            'strip_name': re.compile(r'\s*\[HSP\]'),
            'to_stance': '(~B)HSP',
        },
        {
            'uuid': '3f5cac11-0d1d-42c1-a09e-34f6a66ecdcf',
            'strip_cmd': re.compile(r'\s*\[~B\]'),
            'strip_name': re.compile(r'\s*\[HSP\]'),
            'to_stance': '(~B)HSP',
        },
        {
            'uuid': '5ab8268a-215e-4f04-bf8e-a96f71ac507a',
            'strip_cmd': re.compile(r'\s*\[~D_~B\]'),
            'strip_name': re.compile(r'\s*\[RLX\]'),
            'to_stance': '(~D)RLX; (~B)RLX',
        },
        {
            'uuid': '49e65892-9b94-4709-9f04-ee9f63cf6a53',
            'strip_cmd': re.compile(r'\s*\[~B_~D\]'),
            'strip_name': re.compile(r'\s*\[RLX\]'),
            'to_stance': '(~B)RLX; (~D)RLX',
        },
        {
            'uuid': '7f64d48a-61da-46a7-a8de-d8f58b8776a8',
            'strip_cmd': re.compile(r'\s*\[~D\]'),
            'strip_name': re.compile(r'\s*\[RLX\]'),
            'to_stance': '(~D)RLX',
        },
        {
            'uuid': 'bffcbff7-db7d-4043-ac47-9ee47ec87cc6',
            'strip_cmd': re.compile(r'\s*\[~B\]'),
            'strip_name': re.compile(r'\s*\[HSP\]'),
            'to_stance': '(~B)HSP',
        },
        {
            'uuid': '0a859ca5-1992-4057-a70c-d71fc915cc3c',
            'strip_cmd': re.compile(r'\s*\[~D\]'),
            'strip_name': re.compile(r'\s*\[RLX\]'),
            'to_stance': '(~D)RLX',
        },
        {
            'uuid': '9b85822a-6998-47d6-99b9-a7948eda5324',
            'strip_cmd': re.compile(r'\s*\[~B\]'),
            'strip_name': re.compile(r'\s*\[HSP\]'),
            'to_stance': '(~B)HSP',
        },
        {
            'uuid': 'da4ad8a1-db1f-4376-8397-632562b526f3',
            'strip_cmd': re.compile(r'\s*\[~B\]'),
            'strip_name': re.compile(r'\s*\[HSP\]'),
            'to_stance': '(~B)HSP',
        },
        {
            'uuid': '7b616e47-8ebb-4154-8160-9d5418a3b56c',
            'strip_cmd': re.compile(r'\s*\[~B\]\s*\[~D\]'),
            'strip_name': re.compile(r'\s*\[HSP\]\s*\[RLX\]'),
            'to_stance': '(~B)HSP; (~D)RLX',
        },
        {
            'uuid': 'ce601b9e-9480-4a40-b9a7-373559a41f32',
            'strip_cmd': re.compile(r'\s*\[~D\]'),
            'strip_name': re.compile(r'\s*\[RLX\]'),
            'to_stance': '(~D)RLX',
        },
    ],
    'forest': [
        {
            'uuid': '11cd67de-7112-4c05-acef-561f7c109875',
            'strip_cmd': re.compile(r'\s*\[~D\]'),
            'strip_name': re.compile(r'\s*\[PLD\]'),
            'to_stance': '(~D)PLD',
        },
        {
            'uuid': '751f0015-8ef6-45a1-93b9-a98433ce1588',
            'strip_cmd': re.compile(r'\s*\[~D\]'),
            'strip_name': re.compile(r'\s*\[KND\]'),
            'to_stance': '(~D)KND',
        },
    ],
    'heihachi': [
        {
            'uuid': '5b09c08e-09b3-4913-88a3-14d3bf367ca9',
            'strip_cmd': re.compile(r'\s*\[~u_~d\]'),
            'strip_name': re.compile(r'\s*\[SS\]'),
            'to_stance': '(~u)SS; (~d)SS',
        },
    ],
    'lee': [
        {
            'uuid': '96cf7074-dd8c-4b40-a94e-b20afb923e00',
            'strip_cmd': re.compile(r'\s*\[~3\]'),
            'strip_name': re.compile(r'\s*\[Hit Man Stance\]'),
            'to_stance': '(~3)HMS',
        },
        {
            'uuid': '5b2fca48-dee4-430a-9737-31438e2ffb94',
            'strip_cmd': re.compile(r'\s*\[~3\]'),
            'strip_name': re.compile(r'\s*\[Hit Man Stance\]'),
            'to_stance': '(~3)HMS',
        },
        {
            'uuid': '80993965-6214-47a8-a5aa-f15fddcf516d',
            'strip_cmd': re.compile(r'\s*\[~3\]'),
            'strip_name': re.compile(r'\s*\[Hit Man Stance\]'),
            'to_stance': '(~3)HMS',
        },
    ],
    'hwoarang': [
        {
            'uuid': '1358e51e-baab-4fc9-b47d-2a0df19da788',
            'strip_cmd': re.compile(r'\s*\[F\]'),
            'strip_name': re.compile(r'\s*\[LFF\]'),
            'to_stance': '(~F)LFF',
        },
        {
            'uuid': 'dd80a6d2-b000-466c-ba7b-6261227faef6',
            'strip_cmd': re.compile(r'\s*\[F\]'),
            'strip_name': re.compile(r'\s*\[LFF\]'),
            'to_stance': '(~F)LFF',
        },
        {
            'uuid': 'bef0fe5d-b0fe-4e46-bb14-707a6b87f591',
            'strip_cmd': re.compile(r'\s*\[F\]'),
            'strip_name': re.compile(r'\s*\[LFF\]'),
            'to_stance': '(~F)LFF',
        },
    ],
    'nina': [
        {
            'uuid': 'a4a3c8ec-6ede-43b4-82a0-c27c0bed2c9a',
            'strip_cmd': re.compile(r'\s*\[~U_~D\]'),
            'strip_name': re.compile(r'\s*\[Side Step\]'),
            'to_stance': '(~U)SS; (~D)SS',
        },
        {
            'uuid': 'b982dec9-a4af-40d1-9ef2-a3c1d22aa0bf',
            'strip_cmd': re.compile(r'\s*\[~U_~D\]'),
            'strip_name': re.compile(r'\s*\[Side Step\]'),
            'to_stance': '(~U)SS; (~D)SS',
        },
    ],
    'lei': [
        {
            'uuid': 'ca4b3038-5536-4604-bf2b-b82ab8c8749a',
            'strip_cmd': re.compile(r'\s*\[u_d\]'),
            'strip_name': re.compile(r'\s*\[SNA\]'),
            'to_stance': '(~U)SNA; (~D)SNA',
        },
        {
            'uuid': '492514ee-b8e0-4535-b101-ece3a4b33a47',
            'strip_cmd': re.compile(r'\s*\[u_d\]'),
            'strip_name': re.compile(r'\s*\[DGN\]'),
            'to_stance': '(~U)DGN; (~D)DGN',
        },
        {
            'uuid': 'da2055f8-0385-4fa8-8d5f-12e059367aca',
            'strip_cmd': re.compile(r'\s*\[u_d\]'),
            'strip_name': re.compile(r'\s*\[PAN\]'),
            'to_stance': '(~U)PAN; (~D)PAN',
        },
        {
            'uuid': '30007767-862b-4ddf-a5bd-3f2039dbe303',
            'strip_cmd': re.compile(r'\s*\[u_d\]'),
            'strip_name': re.compile(r'\s*\[TGR\]'),
            'to_stance': '(~U)TGR; (~D)TGR',
        },
        {
            'uuid': '0e5c6c88-5bbf-4f17-989c-3394c26475db',
            'strip_cmd': re.compile(r'\s*\[u_d\]'),
            'strip_name': re.compile(r'\s*\[CRA\]'),
            'to_stance': '(~U)CRA; (~D)CRA',
        },
        {
            'uuid': 'c412201f-6ddf-40f1-9eeb-e4d997199306',
            'strip_cmd': re.compile(r'\s*\[u_d\]'),
            'strip_name': re.compile(r'\s*\[CRA\]'),
            'to_stance': '(~U)CRA; (~D)CRA',
        },
        {
            'uuid': '174b30bb-e33d-474f-9e52-fb638d95beee',
            'strip_cmd': re.compile(r'\s*\[u_d\]'),
            'strip_name': re.compile(r'\s*\[SNA\]'),
            'to_stance': '(~U)SNA; (~D)SNA',
        },
        {
            'uuid': 'c4bfe5c2-fc50-4f5e-9c3d-5c8ec8ba7fc4',
            'strip_cmd': re.compile(r'\s*\[u_d\]'),
            'strip_name': re.compile(r'\s*\[CRA\]'),
            'to_stance': '(~U)CRA; (~D)CRA',
        },
        {
            'uuid': '4a84ffbb-2d77-4046-8cd9-0bf80d105f9f',
            'strip_cmd': re.compile(r'\s*\[u_d\]'),
            'strip_name': re.compile(r'\s*\[SNA\]'),
            'to_stance': '(~U)SNA; (~D)SNA',
        },
        {
            'uuid': '6c82b7d4-2c07-4c94-9843-3be34facb760',
            'strip_cmd': re.compile(r'\s*\[u_d\]'),
            'strip_name': re.compile(r'\s*\[SNA\]'),
            'to_stance': '(~U)SNA; (~D)SNA',
        },
        {
            'uuid': 'a13eb4b8-72a4-45cb-8062-7dde77893db6',
            'strip_cmd': re.compile(r'\s*\[u_d\]'),
            'strip_name': re.compile(r'\s*\[DGN\]'),
            'to_stance': '(~U)DGN; (~D)DGN',
        },
        {
            'uuid': '903f848a-54e6-44bd-8b4f-b57e93ad857f',
            'strip_cmd': re.compile(r'\s*\[u_d\]'),
            'strip_name': re.compile(r'\s*\[PAN\]'),
            'to_stance': '(~U)PAN; (~D)PAN',
        },
        {
            'uuid': '0dbaab4b-b6c6-439a-b6ef-bad1b700ba9e',
            'strip_cmd': re.compile(r'\s*\[u_d\]'),
            'strip_name': re.compile(r'\s*\[TGR\]'),
            'to_stance': '(~U)TGR; (~D)TGR',
        },
        {
            'uuid': '859457ab-cc39-4aef-9726-b489d335c415',
            'strip_cmd': re.compile(r'\s*\[u_d\]'),
            'strip_name': re.compile(r'\s*\[CRA\]'),
            'to_stance': '(~U)CRA; (~D)CRA',
        },
        {
            'uuid': 'a99fb662-c091-428d-b2e7-9749520357d1',
            'strip_cmd': re.compile(r'\s*\[u_d\]'),
            'strip_name': re.compile(r'\s*\[CRA\]'),
            'to_stance': '(~U)CRA; (~D)CRA',
        },
        {
            'uuid': '9d7a2eec-7454-4a25-be57-f2a98c62ad45',
            'strip_cmd': re.compile(r'\s*\[~f\]'),
            'strip_name': re.compile(r'\s*\[DRU\]'),
            'to_stance': '(~F)DRU',
        },
        {
            'uuid': 'b0109d6c-e7a7-4ce3-a61e-e111f396fad9',
            'strip_cmd': re.compile(r'\s*\[~f\]\s*\[~d_~u\]'),
            'strip_name': re.compile(r'\s*\[TGR\]\s*\[DGN\]'),
            'to_stance': '(~F)TGR; (~D)DGN; (~U)DGN',
        },
        {
            'uuid': 'fa7cfc89-03e8-4c62-b819-b4662e3f2e4f',
            'strip_cmd': re.compile(r'\s*\[~f\]\s*\[~d_~u\]'),
            'strip_name': re.compile(r'\s*\[CRA\]\s*\[PAN\]'),
            'to_stance': '(~F)CRA; (~D)PAN; (~U)PAN',
        },
        {
            'uuid': '259c2378-7dc7-4ee1-8eab-738ca43737c3',
            'strip_cmd': re.compile(r'\s*\[~f\]'),
            'strip_name': re.compile(r'\s*\[DRU\]'),
            'to_stance': '(~F)DRU',
        },
        {
            'uuid': '737fd644-160c-4b59-bd70-7042e5a8ee2f',
            'strip_cmd': re.compile(r'\s*\[~F\]'),
            'strip_name': re.compile(r'\s*\[DRU\]'),
            'to_stance': '(~F)DRU',
        },
        {
            'uuid': 'c92d456f-0d26-4aab-ad36-2145cee9d8e0',
            'strip_cmd': re.compile(r'\s*\[~F\]'),
            'strip_name': re.compile(r'\s*\[DRU\]'),
            'to_stance': '(~F)DRU',
        },
        {
            'uuid': '90d6b24a-6358-4cfa-96a0-2c845188c0e2',
            'strip_cmd': re.compile(r'\s*\[~F\]'),
            'strip_name': re.compile(r'\s*\[TGR\]'),
            'to_stance': '(~F)TGR',
        },
        {
            'uuid': '19abb193-8632-4667-8ed2-91fa0c2246ad',
            'strip_cmd': re.compile(r'\s*\[~F\]'),
            'strip_name': re.compile(r'\s*\[TGR\]'),
            'to_stance': '(~F)TGR',
        },
        {
            'uuid': 'f2cc872d-5b2c-49c5-9c5b-cc93b974cd56',
            'strip_cmd': re.compile(r'\s*\[~F\]'),
            'strip_name': re.compile(r'\s*\[SNA\]'),
            'to_stance': '(~F)SNA',
        },
        {
            'uuid': '9bb9abef-5ee0-4534-9c1f-8f347054fe60',
            'strip_cmd': re.compile(r'\s*\[~F\]'),
            'strip_name': re.compile(r'\s*\[DGN\]'),
            'to_stance': '(~F)DGN',
        },
        {
            'uuid': '49903078-2b33-4b0c-8917-3497f1edcc5c',
            'strip_cmd': re.compile(r'\s*\[~F\]'),
            'strip_name': re.compile(r'\s*\[DGN\]'),
            'to_stance': '(~F)DGN',
        },
        {
            'uuid': 'ef28a1dc-5aae-456e-828d-b6d0f124d445',
            'strip_cmd': re.compile(r'\s*\[~F\]'),
            'strip_name': re.compile(r'\s*\[PAN\]'),
            'to_stance': '(~F)PAN',
        },
        {
            'uuid': 'd9110c7e-c3ab-470d-9c68-8e38f863e2af',
            'strip_cmd': re.compile(r'\s*\[d\]'),
            'strip_name': re.compile(r'\s*\[KND\]'),
            'to_stance': '(~D)KND',
        },
        {
            'uuid': '27a50c90-2611-4735-bd22-dd66cda3c26d',
            'strip_cmd': re.compile(r'\s*\[d\]'),
            'strip_name': re.compile(r'\s*\[KND\]'),
            'to_stance': '(~D)KND',
        },
        {
            'uuid': '359b601d-a3e8-4d46-b9b3-51fddcf4bc7e',
            'strip_cmd': re.compile(r'\s*\[d\]'),
            'strip_name': re.compile(r'\s*\[KND\]'),
            'to_stance': '(~D)KND',
        },
        {
            'uuid': '5d38fab1-975c-4345-85d9-186f0dbd11d5',
            'strip_cmd': re.compile(r'\s*\[d\]'),
            'strip_name': re.compile(r'\s*\[KND\]'),
            'to_stance': '(~D)KND',
        },
        {
            'uuid': 'b32879a4-27c9-469a-9ab0-bff47e50f029',
            'strip_cmd': re.compile(r'\s*\[d\]'),
            'strip_name': re.compile(r'\s*\[KND\]'),
            'to_stance': '(~D)KND',
        },
        {
            'uuid': '3551a732-4947-44df-a219-6d7f720bc8fc',
            'strip_cmd': re.compile(r'\s*\[d\]'),
            'strip_name': re.compile(r'\s*\[KND\]'),
            'to_stance': '(~D)KND',
        },
        {
            'uuid': 'eef73e5a-4f90-4050-bd79-c8572cec0794',
            'strip_cmd': re.compile(r'\s*\[~B\]'),
            'strip_name': re.compile(r'\s*\[Art of Phoenix\]'),
            'to_stance': '(~B)AOP',
        },
    ],
    'ling': [
        {
            'uuid': 'fea0d91f-39e3-4926-a9ed-3bc833370e79',
            'strip_cmd': re.compile(r'\s*\[~B\]'),
            'strip_name': re.compile(r'\s*\[RDS\]'),
            'to_stance': '(~B)RDS',
        },
        {
            'uuid': 'eba52a6a-4354-422e-bcd5-c23c11453761',
            'strip_cmd': re.compile(r'\s*\[u_d\]'),
            'strip_name': re.compile(r'\s*\[Roll up or down\]'),
            'to_stance': '(~U)Quick roll to background; (~D)Quick roll to foreground',
        },
    ],
}


def parse_sections(ws):
    sections = []
    row = 1
    while row <= ws.max_row:
        cell = ws.cell(row=row, column=1)
        if cell.value and cell.font.bold:
            next_cell = ws.cell(row=row + 1, column=1) if row + 1 <= ws.max_row else None
            if next_cell and next_cell.value == 'Command' and next_cell.font.bold:
                heading = cell.value
                row += 1

                header_cells = []
                for col in range(1, ws.max_column + 1):
                    val = ws.cell(row=row, column=col).value
                    if val:
                        header_cells.append(val)
                row += 1

                data_rows = []
                while row <= ws.max_row:
                    first_cell = ws.cell(row=row, column=1)
                    if first_cell.value is None and all(
                        ws.cell(row=row, column=c).value is None
                        for c in range(1, len(header_cells) + 1)
                    ):
                        row += 1
                        break
                    if first_cell.font.bold:
                        break

                    row_dict = {}
                    for col_idx, col_name in enumerate(header_cells):
                        val = ws.cell(row=row, column=col_idx + 1).value
                        row_dict[col_name] = val or ''
                    data_rows.append(row_dict)
                    row += 1

                sections.append((heading, data_rows))
            else:
                row += 1
        else:
            row += 1

    return sections


def process_sections(sections, char):
    tag_count = 0
    stance_count = 0

    stance_rules = {}
    for rule in STANCE_RECOVERY.get(char, []):
        stance_rules[rule['uuid']] = rule

    for _, rows in sections:
        for row in rows:
            cmd = str(row.get('Command', ''))

            if TAG_PATTERN.search(cmd):
                tag_count += 1
                row['Taggable'] = 'TRUE'
                row['Command'] = TAG_PATTERN.sub('', cmd).strip()
                alt_cmds = str(row.get('Alt Commands', ''))
                if alt_cmds:
                    row['Alt Commands'] = TAG_PATTERN.sub('', alt_cmds).strip()
                move_name = str(row.get('Move Name', ''))
                row['Move Name'] = NAME_TAG_PATTERN.sub('', move_name).strip()
            else:
                row['Taggable'] = ''

            hit_range = str(row.get('Hit Range', ''))
            if '[!]' in hit_range:
                row['Hit Range'] = hit_range.replace('[!]', '!')

            uuid = str(row.get('UUID', ''))
            if uuid in stance_rules:
                rule = stance_rules[uuid]
                stance_count += 1
                row['Command'] = rule['strip_cmd'].sub('', str(row.get('Command', ''))).strip()
                alt_cmds = str(row.get('Alt Commands', ''))
                if alt_cmds:
                    row['Alt Commands'] = rule['strip_cmd'].sub('', alt_cmds).strip()
                row['Move Name'] = rule['strip_name'].sub('', str(row.get('Move Name', ''))).strip()
                existing = str(row.get('To Stance', ''))
                if existing:
                    row['To Stance'] = existing + '; ' + rule['to_stance']
                else:
                    row['To Stance'] = rule['to_stance']

    return tag_count, stance_count


OUTPUT_COLUMNS = [
    'Command', 'Alt Commands', 'Move Name', 'To Stance',
    'Speed', 'Block Adv', 'Hit Adv', 'Counter Hit Adv',
    'Damage', 'Damage Sum', 'Hit Range',
    'Throw Type', 'Throw Escape', 'Properties', 'Notes',
    'Taggable',
    'UUID', 'ML UUID', 'Parent UUID', 'Unmatched',
]


HIDDEN_COLUMNS = set()


def autofit_columns(ws, header_row_numbers):
    for col in ws.columns:
        col_letter = col[0].column_letter
        header_value = None
        for cell in col:
            if cell.row in header_row_numbers and cell.value:
                header_value = cell.value
                break
        if header_value in HIDDEN_COLUMNS:
            ws.column_dimensions[col_letter].width = 0
            ws.column_dimensions[col_letter].hidden = True
            continue
        max_len = 0
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max_len + 2


def find_characters():
    files = glob.glob('sources/*_step6.xlsx')
    return sorted(
        os.path.basename(f).replace('_step6.xlsx', '')
        for f in files
        if not os.path.basename(f).startswith('~$')
    )


def main():
    characters = find_characters()
    if not characters:
        print("No step6 files found in sources/")
        return
    print(f"Found characters: {characters}")

    for char in characters:
        input_path = f"sources/{char}_step6.xlsx"
        output_path = f"sources/{char}_step6b.xlsx"

        print(f"\n=== {char.capitalize()} ===")
        wb_in = load_workbook(input_path)
        ws_in = wb_in.active

        sections = parse_sections(ws_in)
        tag_count, stance_count = process_sections(sections, char)

        print(f"  {tag_count} rows tagged, {stance_count} stance recoveries")

        wb_out = Workbook()
        ws = wb_out.active
        ws.title = "Merged"
        row_num = 1
        header_row_numbers = set()

        for section_name, rows in sections:
            cell = ws.cell(row=row_num, column=1, value=section_name)
            cell.font = Font(bold=True)
            row_num += 1

            for col, h in enumerate(OUTPUT_COLUMNS, 1):
                cell = ws.cell(row=row_num, column=col, value=h)
                cell.font = Font(bold=True)
            header_row_numbers.add(row_num)
            row_num += 1

            for row_dict in rows:
                for col, col_name in enumerate(OUTPUT_COLUMNS, 1):
                    val = row_dict.get(col_name, '')
                    if val:
                        ws.cell(row=row_num, column=col, value=val)
                row_num += 1

            row_num += 1

        autofit_columns(ws, header_row_numbers)
        wb_out.save(output_path)
        print(f"  Written to {output_path}")

    print(f"\nDone. Processed {len(characters)} characters.")


if __name__ == '__main__':
    main()
