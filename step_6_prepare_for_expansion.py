#!/usr/bin/env python3
"""
Step 6: Prepare data for continuation expansion in step 7.

Reads sources/<character>_step5b.xlsx and fixes Move Name values that
would otherwise expand incorrectly.

Output: sources/<character>_step6.xlsx

Usage:
    python3 step_6_prepare_for_expansion.py
"""

import glob
import os

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font


CORRECTIONS = {
    'gunjack': {
        'modify': [],
        'insert_after': [],
        'delete': [],
        'replace': [
            {
                'uuid': 'ab11e12d-c483-4b93-9879-06d76d1fca71',
                'rows': [
                    {
                        'Command': 'db+1,1,1,2',
                        'Move Name': 'Machine Gun - Megaton Punch',
                        'Speed': '12',
                        'Block Adv': '-6 -6 -11 -18',
                        'Hit Adv': '+5 +5 0 KD',
                        'Counter Hit Adv': '+5 +5 0 KD',
                        'Damage': '15,15,15,40',
                        'Damage Sum': '85',
                        'Hit Range': 'lllm',
                        'UUID': '3dfc05af-5516-438c-9613-534804a1c3bb',
                    },
                    {
                        'Command': 'db+1,1,1,1,2',
                        'Move Name': 'Machine Gun - Megaton Punch',
                        'Speed': '12',
                        'Block Adv': '-6 -6 -6 -11 -18',
                        'Hit Adv': '+5 +5 +5 0 KD',
                        'Counter Hit Adv': '+5 +5 +5 0 KD',
                        'Damage': '15,15,15,15,40',
                        'Damage Sum': '100',
                        'Hit Range': 'llllm',
                        'UUID': 'cb9ca34f-2dcd-4444-bb8a-409bc1f3ce79',
                    },
                    {
                        'Command': 'db+1,1,1,1,1,2',
                        'Move Name': 'Machine Gun - Megaton Punch',
                        'Speed': '12',
                        'Block Adv': '-6 -6 -6 -6 -11 -18',
                        'Hit Adv': '+5 +5 +5 +5 0 KD',
                        'Counter Hit Adv': '+5 +5 +5 +5 0 KD',
                        'Damage': '15,15,15,15,15,40',
                        'Damage Sum': '115',
                        'Hit Range': 'lllllm',
                        'UUID': '3fc4e6c7-b2f3-4a63-944b-235077793821',
                    },
                ],
            },
        ],
    },
    'jack2': {
        'modify': [],
        'insert_after': [],
        'delete': [],
        'replace': [
            {
                'uuid': '4cf7be1b-11d0-4c05-8424-e597df98603b',
                'rows': [
                    {
                        'Command': 'db+1,1,1,2',
                        'Move Name': 'Machine Gun - Megaton Punch',
                        'Speed': '12',
                        'Block Adv': '-6 -6 -11 -16',
                        'Hit Adv': '+5 +5 0 KD',
                        'Counter Hit Adv': '+5 +5 0 KD',
                        'Damage': '12,15,15,40',
                        'Damage Sum': '82',
                        'Hit Range': 'lllm',
                        'UUID': 'dc5a29f5-7ecb-4a43-967e-7b5eadf20512',
                    },
                    {
                        'Command': 'db+1,1,1,1,2',
                        'Move Name': 'Machine Gun - Megaton Punch',
                        'Speed': '12',
                        'Block Adv': '-6 -6 -6 -11 -16',
                        'Hit Adv': '+5 +5 +5 0 KD',
                        'Counter Hit Adv': '+5 +5 +5 0 KD',
                        'Damage': '12,15,15,15,40',
                        'Damage Sum': '97',
                        'Hit Range': 'llllm',
                        'UUID': 'b3cad2bc-95a3-4d7c-9abb-0335ed57fb7d',
                    },
                    {
                        'Command': 'db+1,1,1,1,1,2',
                        'Move Name': 'Machine Gun - Megaton Punch',
                        'Speed': '12',
                        'Block Adv': '-6 -6 -6 -6 -11 -16',
                        'Hit Adv': '+5 +5 +5 +5 0 KD',
                        'Counter Hit Adv': '+5 +5 +5 +5 0 KD',
                        'Damage': '12,15,15,15,15,40',
                        'Damage Sum': '112',
                        'Hit Range': 'lllllm',
                        'UUID': '3a6ce42c-3b05-4641-9152-428a8ed9fd0f',
                    },
                ],
            },
        ],
    },
    'lee': {
        'modify': [],
        'insert_after': [],
        'delete': [],
        'replace': [
            {
                'uuid': 'a1226d52-19ca-4e28-a360-74af905c7295',
                'rows': [
                    {
                        'Command': 'b+3~3',
                        'Move Name': 'Feint Mist Wolf',
                        'Speed': '14',
                        'Block Adv': '-8',
                        'Hit Adv': '+3',
                        'Counter Hit Adv': '+3',
                        'Damage': '18',
                        'Damage Sum': '18',
                        'Hit Range': 'h',
                        'UUID': 'a1226d52-19ca-4e28-a360-74af905c7295',
                        'ML UUID': 'c92e1826-7388-447a-858b-7a3d0c9ce2f3',
                    },
                    {
                        'Command': 'b+3~3:4',
                        'Move Name': 'Mist Trap',
                        'Damage': '33',
                        'Damage Sum': '33',
                        'Notes': 'If the High Kick is blocked, tapping 4 at the exact frame will cause the opponent to grab Lee\'s foot. Lee will auto reverse.',
                        'UUID': '6d339f6a-39b1-4eb1-bcd3-ded0f8935a0c',
                        'ML UUID': 'c92e1826-7388-447a-858b-7a3d0c9ce2f3',
                        'Parent UUID': 'a1226d52-19ca-4e28-a360-74af905c7295',
                    },
                ],
            },
        ],
    },
    'julia': {
        'modify': [
            {
                'uuid': '697b57ac-f4e1-4b88-af3c-1c93b7f5ef92',
                'set': {'Move Name': '– Bow - Arrow'},
            },
            {
                'uuid': '15cd3b1a-c236-4443-abd8-0bb0c83297ea',
                'set': {'Move Name': '– Bow - Arrow'},
            },
            {
                'uuid': 'ecf9ec9c-b22c-4f6a-9f5c-8f0b3d51b517',
                'set': {'Move Name': 'Club Fist - Bow - Arrow'},
            },
            {
                'uuid': '3308856d-d15f-45c7-8281-db3de8f037ce',
                'set': {'Move Name': '– Club Fist - Bow - Arrow'},
            },
            {
                'uuid': 'f4f51a05-dcc8-4a5f-a2e8-a423c1aceae5',
                'set': {'Move Name': 'Bow - Arrow'},
            },
        ],
        'insert_after': [],
        'delete': [],
    },
    'lei': {
        'modify': [
            {
                'uuid': 'b32879a4-27c9-469a-9ab0-bff47e50f029',
                'set': {
                    'Command': 'b+4 [d]',
                    'Move Name': 'Tornado Kick [KND]',
                    'Parent UUID': '',
                },
            },
            {
                'uuid': '3551a732-4947-44df-a219-6d7f720bc8fc',
                'set': {
                    'Command': 'b+4,U [d]',
                    'Move Name': 'Triple Tornado [KND]',
                    'Parent UUID': '',
                },
            },
        ],
        'insert_after': [],
        'delete': [
            'a12ecda6-d3e8-4886-92f4-5d3e68ede144',
            '67124b94-0dbe-46d2-85c5-9d5009dcb5d7',
        ],
        'move_to_new_section': {
            'section_name': 'Art of Phoenix Illusion',
            'uuids': [
                'b32879a4-27c9-469a-9ab0-bff47e50f029',
                '3551a732-4947-44df-a219-6d7f720bc8fc',
            ],
            'append_rows': [
                {
                    'Command': '4',
                    'Move Name': 'Phoenix Strike',
                    'Speed': '68',
                    'Block Adv': '',
                    'Hit Adv': 'KD',
                    'Counter Hit Adv': 'KD',
                    'Damage': '90',
                    'Damage Sum': '90',
                    'Hit Range': '!',
                    'UUID': 'dbe04986-7a15-4889-97d3-99ceb7c4607c',
                },
                {
                    'Command': '3,4',
                    'Move Name': 'Hopping Phoenix Kick - Phoenix Strike',
                    'Speed': '14',
                    'Block Adv': '-26',
                    'Hit Adv': 'KD KD',
                    'Counter Hit Adv': 'KD KD',
                    'Damage': '15,90',
                    'Damage Sum': '105',
                    'Hit Range': 'm!',
                    'UUID': '733cdbab-9b5a-4472-97d9-6b5dd8036bcb',
                },
                {
                    'Command': '3,3,4',
                    'Move Name': 'Hopping Phoenix Kicks - Phoenix Strike',
                    'Speed': '14',
                    'Block Adv': '-26 -26',
                    'Hit Adv': 'KD KD KD',
                    'Counter Hit Adv': 'KD KD KD',
                    'Damage': '15,15,90',
                    'Damage Sum': '120',
                    'Hit Range': 'mm!',
                    'UUID': '8fdac8ab-1cea-4bfb-a4df-bc746259e97d',
                },
                {
                    'Command': '3,3,3,4',
                    'Move Name': 'Hopping Phoenix Kicks - Phoenix Strike',
                    'Speed': '14',
                    'Block Adv': '-26 -26 -26',
                    'Hit Adv': 'KD KD KD KD',
                    'Counter Hit Adv': 'KD KD KD KD',
                    'Damage': '15,15,15,90',
                    'Damage Sum': '135',
                    'Hit Range': 'mmm!',
                    'UUID': '3412743d-59da-414c-9027-7931cb5f4c87',
                },
                {
                    'Command': '3,3,3,3,4',
                    'Move Name': 'Hopping Phoenix Kicks - Phoenix Strike',
                    'Speed': '14',
                    'Block Adv': '-26 -26 -26 -26',
                    'Hit Adv': 'KD KD KD KD KD',
                    'Counter Hit Adv': 'KD KD KD KD KD',
                    'Damage': '15,15,15,15,90',
                    'Damage Sum': '150',
                    'Hit Range': 'mmmm!',
                    'UUID': 'c27947cc-5587-46f2-a122-7925bee8fb3a',
                },
            ],
        },
    },
}


def parse_sections(ws):
    """Parse a merged worksheet into sections."""
    sections = []
    row = 1
    while row <= ws.max_row:
        cell = ws.cell(row=row, column=1)
        if cell.value and cell.font.bold:
            next_cell = ws.cell(row=row + 1, column=1) if row + 1 <= ws.max_row else None
            if next_cell and next_cell.value == 'Command' and next_cell.font.bold:
                heading = cell.value
                row += 1

                header_cells = []
                for col in range(1, ws.max_column + 1):
                    val = ws.cell(row=row, column=col).value
                    if val:
                        header_cells.append(val)
                row += 1

                data_rows = []
                while row <= ws.max_row:
                    first_cell = ws.cell(row=row, column=1)
                    if first_cell.value is None and all(
                        ws.cell(row=row, column=c).value is None
                        for c in range(1, len(header_cells) + 1)
                    ):
                        row += 1
                        break
                    if first_cell.font.bold:
                        break

                    row_dict = {}
                    for col_idx, col_name in enumerate(header_cells):
                        val = ws.cell(row=row, column=col_idx + 1).value
                        row_dict[col_name] = val or ''
                    data_rows.append(row_dict)
                    row += 1

                sections.append((heading, header_cells, data_rows))
            else:
                row += 1
        else:
            row += 1

    return sections


def apply_corrections(sections, corrections):
    """Apply modify/insert/delete/replace/move corrections to parsed sections."""
    uuid_col = 'UUID'

    for mod in corrections.get('modify', []):
        target_uuid = mod['uuid']
        found = False
        for _, _, rows in sections:
            for row in rows:
                if row.get(uuid_col) == target_uuid:
                    for col_name, value in mod['set'].items():
                        row[col_name] = value
                    found = True
                    break
            if found:
                break
        if not found:
            print(f"  WARNING: UUID {target_uuid} not found for modify")

    for ins in corrections.get('insert_after', []):
        after_uuid = ins['after_uuid']
        new_row = ins['row']
        found = False
        for _, header_cells, rows in sections:
            for i, row in enumerate(rows):
                if row.get(uuid_col) == after_uuid:
                    row_dict = {col: '' for col in header_cells}
                    row_dict.update(new_row)
                    rows.insert(i + 1, row_dict)
                    found = True
                    break
            if found:
                break
        if not found:
            print(f"  WARNING: UUID {after_uuid} not found for insert_after")

    for delete_uuid in corrections.get('delete', []):
        found = False
        for _, _, rows in sections:
            for i, row in enumerate(rows):
                if row.get(uuid_col) == delete_uuid:
                    rows.pop(i)
                    found = True
                    break
            if found:
                break
        if not found:
            print(f"  WARNING: UUID {delete_uuid} not found for delete")

    for repl in corrections.get('replace', []):
        target_uuid = repl['uuid']
        found = False
        for _, header_cells, rows in sections:
            for i, row in enumerate(rows):
                if row.get(uuid_col) == target_uuid:
                    rows.pop(i)
                    for j, new_row in enumerate(repl['rows']):
                        row_dict = {col: '' for col in header_cells}
                        row_dict.update(new_row)
                        rows.insert(i + j, row_dict)
                    found = True
                    break
            if found:
                break
        if not found:
            print(f"  WARNING: UUID {target_uuid} not found for replace")

    move = corrections.get('move_to_new_section')
    if move:
        section_name = move['section_name']
        target_uuids = set(move['uuids'])
        header_cells = None
        moved_rows = []
        for _, hc, rows in sections:
            if header_cells is None:
                header_cells = hc
            to_remove = []
            for i, row in enumerate(rows):
                if row.get(uuid_col) in target_uuids:
                    to_remove.append(i)
                    moved_rows.append(row)
            for i in reversed(to_remove):
                rows.pop(i)

        for new_row in move.get('append_rows', []):
            row_dict = {col: '' for col in header_cells}
            row_dict.update(new_row)
            moved_rows.append(row_dict)

        sections.append((section_name, header_cells, moved_rows))

    return sections


def autofit_columns(ws):
    """Auto-fit all column widths based on cell content."""
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max_len + 2


def write_sections(sections, output_path):
    """Write sections to a new xlsx file."""
    wb_out = Workbook()
    ws = wb_out.active
    ws.title = "Merged"
    row_num = 1

    for section_name, header_cells, rows in sections:
        cell = ws.cell(row=row_num, column=1, value=section_name)
        cell.font = Font(bold=True)
        row_num += 1

        for col, h in enumerate(header_cells, 1):
            cell = ws.cell(row=row_num, column=col, value=h)
            cell.font = Font(bold=True)
        row_num += 1

        for row_dict in rows:
            for col, col_name in enumerate(header_cells, 1):
                val = row_dict.get(col_name, '')
                if val:
                    ws.cell(row=row_num, column=col, value=val)
            row_num += 1

        row_num += 1

    autofit_columns(ws)
    wb_out.save(output_path)


def find_characters():
    """Find all characters that have step5b xlsx files."""
    files = glob.glob('sources/*_step5b.xlsx')
    return sorted(
        os.path.basename(f).replace('_step5b.xlsx', '')
        for f in files
        if not os.path.basename(f).startswith('~$')
    )


def main():
    characters = find_characters()
    if not characters:
        print("No step5b files found in sources/")
        return
    print(f"Found characters: {characters}")

    for char in characters:
        input_path = f"sources/{char}_step5b.xlsx"
        output_path = f"sources/{char}_step6.xlsx"

        print(f"\n=== {char.capitalize()} ===")
        wb_in = load_workbook(input_path)
        ws_in = wb_in.active

        sections = parse_sections(ws_in)

        corrections = CORRECTIONS.get(char, {'modify': [], 'insert_after': [], 'delete': []})
        sections = apply_corrections(sections, corrections)

        write_sections(sections, output_path)
        print(f"  Written to {output_path}")

    print(f"\nDone. Processed {len(characters)} characters.")


if __name__ == '__main__':
    main()
