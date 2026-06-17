#!/usr/bin/env python3
"""
Step 9: Split multi-hit rows into separate rows, one per press.

Reads sources/<character>_step8.xlsx. Rows with multiple space-separated
Block Adv Full values are split into progressive rows where each row
represents the combo up to that press.

Output: sources/<character>_step9.xlsx

Usage:
    python3 step_9_split_hits.py
"""

import glob
import os
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font


DIRECTIONS = {'f', 'b', 'd', 'u', 'df', 'db', 'uf', 'ub', 'N', 'F', 'B', 'D', 'U',
              'DF', 'DB', 'UF', 'UB', 'FC'}

DONT_SPLIT = {
    ('jin', 'd+3+4'),
    ('ling', 'd+1'),
    ('ling', 'D+1'),
    ('ling', 'u+1+2'),
    ('ling', 'SS 4 [~B]'),
}

IGNORE = {('ling', 'FC,DF+2')}


def token_has_button(tok):
    return bool(re.search(r'[1-4]', tok))


def add_spaces(cmd):
    """Insert spaces between presses in a command string."""
    trailing = ''
    m = re.search(r'(\s+-\s*(\[.*?\])?\s*)$', cmd)
    if m:
        trailing = m.group(0)
        cmd_clean = cmd[:m.start()]
    else:
        cmd_clean = cmd

    result = []
    i = 0
    while i < len(cmd_clean):
        ch = cmd_clean[i]
        if ch == '<':
            result.append(' ')
            result.append(ch)
        elif ch == '~':
            rest = cmd_clean[i+1:]
            next_sep = re.search(r'[,<~]', rest)
            next_part = rest[:next_sep.start()] if next_sep else rest
            if token_has_button(next_part):
                result.append(' ')
            result.append(ch)
        elif ch == ',':
            before = ''.join(result).split(' ')[-1] if result else ''
            rest = cmd_clean[i+1:]
            next_sep = re.search(r'[,<~]', rest)
            next_part = rest[:next_sep.start()] if next_sep else rest

            prev_has_btn = token_has_button(before)
            prev_ends_dir = False
            if '~' not in before:
                parts = before.split('<')
                last = parts[-1]
                if last in DIRECTIONS:
                    prev_ends_dir = True

            curr_is_dir_btn = bool(re.match(r'^[fbudFBUDN][fbudFBUDN]*\+[1-4]', next_part))
            curr_is_motion = '~' in next_part
            curr_has_btn = token_has_button(next_part)
            prev_is_pure_dir = not token_has_button(before)
            curr_is_pure_dir = not curr_has_btn

            if prev_is_pure_dir and curr_is_pure_dir:
                result.append(ch)
            elif prev_is_pure_dir and curr_is_dir_btn:
                result.append(ch)
            elif prev_is_pure_dir and curr_is_motion:
                result.append(ch)
            elif prev_is_pure_dir and curr_has_btn and not curr_is_dir_btn:
                result.append(' ')
                result.append(ch)
            elif prev_has_btn and prev_ends_dir:
                result.append(ch)
            elif prev_has_btn and not prev_ends_dir and curr_has_btn:
                result.append(' ')
                result.append(ch)
            elif prev_has_btn and curr_is_pure_dir:
                result.append(ch)
            else:
                result.append(ch)
        else:
            result.append(ch)
        i += 1

    final = ''.join(result).strip() + trailing
    final = re.sub(r'  +', ' ', final)
    return final


def get_press_tokens(spaced_cmd):
    """Split spaced command into press tokens, handling stance prefix and trailing."""
    trailing = ''
    m = re.search(r'(\s+-\s*(\[.*?\])?\s*)$', spaced_cmd)
    if m:
        trailing = m.group(0)
        clean = spaced_cmd[:m.start()]
    else:
        clean = spaced_cmd

    prefix = ''
    prefix_match = re.match(r'^(WS|SS|RDS|FC)\s+', clean)
    if prefix_match:
        prefix = prefix_match.group(0)
        clean = clean[prefix_match.end():]

    tokens = clean.split(' ')
    return prefix, tokens, trailing


def build_progressive_commands(prefix, tokens, trailing):
    """Build progressive command list from press tokens."""
    commands = []
    for i in range(1, len(tokens) + 1):
        cmd = prefix + ''.join(tokens[:i])
        if i == len(tokens) and trailing:
            cmd += trailing
        commands.append(cmd)
    return commands


def assign_values_to_presses(values, press_count):
    """Assign block/hit/counter values (space-separated) to presses.

    When there are more values than presses (multi-hit press case),
    group 'x' with the following value. When counts match exactly,
    assign one value per press without grouping.
    """
    if len(values) == press_count:
        return list(values)

    assignments = []
    i = 0
    while i < len(values) and len(assignments) < press_count:
        if values[i] == 'x':
            if i + 1 < len(values):
                assignments.append(values[i] + ' ' + values[i + 1])
                i += 2
            else:
                assignments.append(values[i])
                i += 1
        else:
            assignments.append(values[i])
            i += 1
    while len(assignments) < press_count:
        assignments.append('')
    return assignments


def hits_per_press_from_blocks(block_values, press_count):
    """Calculate how many hits each press produces based on x markers."""
    if len(block_values) == press_count:
        return [1] * press_count
    hits = []
    i = 0
    for _ in range(press_count):
        if i < len(block_values) and block_values[i] == 'x':
            hits.append(2)
            i += 2
        else:
            hits.append(1)
            i += 1
    return hits


def split_range(range_full, press_count, block_values):
    """Split Hit Range Full into per-press groups."""
    if not range_full:
        return [''] * press_count

    hits_per_press = hits_per_press_from_blocks(block_values, press_count)

    if sum(hits_per_press) == len(range_full):
        result = []
        pos = 0
        for count in hits_per_press:
            result.append(range_full[pos:pos + count])
            pos += count
        return result

    if len(range_full) == len(block_values):
        if len(block_values) == press_count:
            return list(range_full)
        return assign_values_to_presses(list(range_full), press_count)

    groups = re.findall(r'[^-]-*', range_full)
    if len(groups) == press_count:
        return groups

    if len(groups) == len(block_values):
        result = []
        gi = 0
        for count in hits_per_press:
            chunk = ''.join(groups[gi:gi + count])
            result.append(chunk)
            gi += count
        return result

    return [''] * (press_count - 1) + [range_full]


def split_damage(damage_full, press_count, block_values):
    """Split Damage Full (comma-separated) into per-press groups."""
    if not damage_full:
        return [''] * press_count

    parts = [x.strip() for x in str(damage_full).split(',')]

    hits_per_press = hits_per_press_from_blocks(block_values, press_count)

    if len(parts) == press_count:
        return parts

    if len(parts) == len(block_values):
        result = []
        pi = 0
        for count in hits_per_press:
            chunk = ','.join(parts[pi:pi + count])
            result.append(chunk)
            pi += count
        return result

    has_dashes = any(p == '-' for p in parts)
    if has_dashes:
        groups = []
        current = []
        for p in parts:
            current.append(p)
            if p != '-':
                groups.append(','.join(current))
                current = []
        if current:
            if groups:
                groups[-1] += ',' + ','.join(current)
            else:
                groups.append(','.join(current))
        if len(groups) == press_count:
            return groups
        if len(groups) == len(block_values):
            result = []
            gi = 0
            for count in hits_per_press:
                chunk = ','.join(groups[gi:gi + count])
                result.append(chunk)
                gi += count
            return result

    return [''] * (press_count - 1) + [damage_full]


def split_row(row, char):
    """Split a multi-hit row into multiple rows. Returns list of row dicts."""
    cmd_full = str(row.get('Command Full', ''))
    block_full = str(row.get('Block Adv Full', ''))
    hit_full = str(row.get('Hit Adv Full', ''))
    ch_full = str(row.get('Counter Hit Adv Full', ''))
    damage_full = str(row.get('Damage Full', ''))
    range_full = str(row.get('Hit Range Full', ''))

    block_values = block_full.split()
    hit_count = len(block_values)

    if hit_count <= 1:
        return [row]

    if (char, cmd_full) in DONT_SPLIT:
        return [row]

    if (char, cmd_full) in IGNORE:
        return [row]

    spaced = add_spaces(cmd_full)
    prefix, tokens, trailing = get_press_tokens(spaced)
    press_count = len(tokens)

    if press_count > hit_count:
        return [row]

    block_assignments = assign_values_to_presses(block_values, press_count)

    hit_values = hit_full.split() if hit_full else []
    hit_assignments = assign_values_to_presses(hit_values, press_count) if hit_values else [''] * press_count

    ch_values = ch_full.split() if ch_full else []
    ch_assignments = assign_values_to_presses(ch_values, press_count) if ch_values else [''] * press_count

    progressive_cmds = build_progressive_commands(prefix, tokens, trailing)
    range_assignments = split_range(range_full, press_count, block_values)
    damage_assignments = split_damage(damage_full, press_count, block_values)

    result_rows = []
    for i in range(press_count):
        new_row = dict(row)
        new_row['Command Full'] = progressive_cmds[i]
        new_row['Block Adv Full'] = block_assignments[i]
        new_row['Hit Adv Full'] = hit_assignments[i]
        new_row['Counter Hit Adv Full'] = ch_assignments[i]
        new_row['Damage Full'] = damage_assignments[i]
        new_row['Hit Range Full'] = range_assignments[i]
        new_row['Hit Index'] = str(i + 1)
        result_rows.append(new_row)

    return result_rows


def parse_sections(ws):
    """Parse a worksheet into sections."""
    sections = []
    row = 1
    while row <= ws.max_row:
        cell = ws.cell(row=row, column=1)
        if cell.value and cell.font.bold:
            next_cell = ws.cell(row=row + 1, column=1) if row + 1 <= ws.max_row else None
            if next_cell and next_cell.value == 'Command' and next_cell.font.bold:
                heading = cell.value
                row += 1

                header_cells = []
                for col in range(1, ws.max_column + 1):
                    val = ws.cell(row=row, column=col).value
                    if val:
                        header_cells.append(val)
                row += 1

                data_rows = []
                while row <= ws.max_row:
                    first_cell = ws.cell(row=row, column=1)
                    if first_cell.value is None and all(
                        ws.cell(row=row, column=c).value is None
                        for c in range(1, len(header_cells) + 1)
                    ):
                        row += 1
                        break
                    if first_cell.font.bold:
                        break

                    row_dict = {}
                    for col_idx, col_name in enumerate(header_cells):
                        val = ws.cell(row=row, column=col_idx + 1).value
                        row_dict[col_name] = val or ''
                    data_rows.append(row_dict)
                    row += 1

                sections.append((heading, data_rows))
            else:
                row += 1
        else:
            row += 1

    return sections


def process_sections(sections, char):
    """Split multi-hit rows. Returns (new_sections, split_count)."""
    split_count = 0
    new_sections = []

    for section_name, rows in sections:
        new_rows = []
        for row in rows:
            result = split_row(row, char)
            if len(result) > 1:
                split_count += 1
            new_rows.extend(result)
        new_sections.append((section_name, new_rows))

    return new_sections, split_count


OUTPUT_COLUMNS = [
    'Command', 'Command Full', 'Alt Commands', 'Move Name', 'Move Name Full', 'Stance',
    'Hit', 'Block Adv', 'Block Adv Full', 'Hit Adv', 'Hit Adv Full',
    'Counter Hit Adv', 'Counter Hit Adv Full', 'Damage', 'Damage Sum', 'Damage Full',
    'Hit Range', 'Hit Range Full', 'Throw Type', 'Throw Escape', 'Properties', 'Notes',
    'Taggable', 'Hit Index', 'UUID', 'ML UUID', 'Parent UUID', 'Unmatched',
]

HIDDEN_COLUMNS = {'Command', 'Move Name', 'Damage', 'Damage Sum', 'Hit Range',
                  'Block Adv', 'Hit Adv', 'Counter Hit Adv'}


def autofit_columns(ws, header_row_numbers):
    """Auto-fit column widths. Hidden columns get width 0."""
    for col in ws.columns:
        col_letter = col[0].column_letter
        header_value = None
        for cell in col:
            if cell.row in header_row_numbers and cell.value:
                header_value = cell.value
                break
        if header_value in HIDDEN_COLUMNS:
            ws.column_dimensions[col_letter].width = 0
            ws.column_dimensions[col_letter].hidden = True
            continue
        max_len = 0
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max_len + 2


def find_characters():
    """Find all characters that have step8 xlsx files."""
    files = glob.glob('sources/*_step8.xlsx')
    return sorted(
        os.path.basename(f).replace('_step8.xlsx', '')
        for f in files
        if not os.path.basename(f).startswith('~$')
    )


def main():
    characters = find_characters()
    if not characters:
        print("No step8 files found in sources/")
        return
    print(f"Found characters: {characters}")

    for char in characters:
        input_path = f"sources/{char}_step8.xlsx"
        output_path = f"sources/{char}_step9.xlsx"

        print(f"\n=== {char.capitalize()} ===")
        wb_in = load_workbook(input_path)
        ws_in = wb_in.active

        sections = parse_sections(ws_in)
        new_sections, split_count = process_sections(sections, char)

        total_rows = sum(len(rows) for _, rows in new_sections)
        original_rows = sum(len(rows) for _, rows in sections)
        print(f"  {split_count} rows split ({original_rows} → {total_rows} rows)")

        wb_out = Workbook()
        ws = wb_out.active
        ws.title = "Merged"
        row_num = 1
        header_row_numbers = set()

        for section_name, rows in new_sections:
            cell = ws.cell(row=row_num, column=1, value=section_name)
            cell.font = Font(bold=True)
            row_num += 1

            for col, h in enumerate(OUTPUT_COLUMNS, 1):
                cell = ws.cell(row=row_num, column=col, value=h)
                cell.font = Font(bold=True)
            header_row_numbers.add(row_num)
            row_num += 1

            for row_dict in rows:
                for col, col_name in enumerate(OUTPUT_COLUMNS, 1):
                    val = row_dict.get(col_name, '')
                    if val:
                        ws.cell(row=row_num, column=col, value=val)
                row_num += 1

            row_num += 1

        autofit_columns(ws, header_row_numbers)
        wb_out.save(output_path)
        print(f"  Written to {output_path}")

    print(f"\nDone. Processed {len(characters)} characters.")


if __name__ == '__main__':
    main()
