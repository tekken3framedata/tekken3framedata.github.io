#!/usr/bin/env python3
"""
Step 1b: Fix Anna source data issues.

Reads sources/anna_step1.xlsx and fixes the 1,2 / d/f+1,2 / BT (1_2),2 problem:
the source has d/f+1,2 and BT (1_2),2 placed before 1,2 as if they were string
starters for the continuation moves (= 3, = 4, etc.). In reality they are
independent moves that happen to share the same string extensions. This script
moves them to their correct position (after 1,4) and removes their string hit
variants.

Removes from String Hit Arts:
  - d/f+1,2,1,2,3,3,2,1,2,4
  - (BT 1_2),2,1,2,3,3,2,1,2,4

Output: sources/anna_step1b.xlsx

Usage:
    python3 step_1b_anna_fixes.py
"""

import os

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font


UUIDS_TO_REMOVE = {
    'Frame Data': [],
    'Movelist': [
        '3868d410-2e89-48d9-953a-b6d47bab49c4',
        '2f43bade-2f95-4531-9ea3-4cdbb2bdd0bd',
    ],
}

UUIDS_TO_MOVE = {
    'Frame Data': {
        'after_uuid': '3db16d2c-f04e-40cb-a74c-ecbfde854d6d',
        'move_uuids': [
            'd0c5a7da-b497-4286-a857-fd8930564a52',
            '2f8871a4-5b86-4d12-b379-301dfb545e64',
        ],
    },
    'Movelist': {
        'after_uuid': 'eb2690e4-5fed-4120-8ae5-c870948f502d',
        'move_uuids': [
            '38c084ac-b201-4176-b5c8-6332f4f011fb',
            '88b77f98-207e-4a96-bf95-31b06c11f4ac',
        ],
    },
}


def get_row_uuid(ws, row):
    """Return the UUID found in any cell of the row, or None."""
    for col in range(1, ws.max_column + 1):
        val = ws.cell(row=row, column=col).value
        if val and len(str(val).strip()) == 36 and '-' in str(val):
            return str(val).strip()
    return None


def read_all_rows(ws):
    """Read all rows as (uuid, [values], [bold_flags]) tuples."""
    rows = []
    for row in range(1, ws.max_row + 1):
        values = []
        bolds = []
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=row, column=col)
            values.append(cell.value)
            bolds.append(bool(cell.font.bold))
        uuid = get_row_uuid(ws, row)
        rows.append((uuid, values, bolds))
    return rows


def process_sheet(ws_in, ws_out, uuids_to_remove, move_spec):
    """Process sheet: remove rows, move rows to new position."""
    rows = read_all_rows(ws_in)
    remove_set = set(uuids_to_remove)
    move_uuids = move_spec['move_uuids'] if move_spec else []
    after_uuid = move_spec['after_uuid'] if move_spec else None
    move_set = set(move_uuids)

    moved_rows = [r for r in rows if r[0] in move_set]
    remaining = [r for r in rows if r[0] not in remove_set and r[0] not in move_set]

    result = []
    for r in remaining:
        result.append(r)
        if r[0] == after_uuid:
            for mr in moved_rows:
                result.append(mr)

    out_row = 0
    for uuid, values, bolds in result:
        out_row += 1
        for col, (val, bold) in enumerate(zip(values, bolds), 1):
            cell = ws_out.cell(row=out_row, column=col, value=val)
            if bold:
                cell.font = Font(bold=True)

    return len([r for r in rows if r[0] in remove_set]), len(moved_rows)


def main():
    input_path = 'sources/anna_step1.xlsx'
    output_path = 'sources/anna_step1b.xlsx'

    if not os.path.exists(input_path):
        print(f"Input file not found: {input_path}")
        return

    print(f"Reading {input_path}...")
    wb_in = load_workbook(input_path)
    wb_out = Workbook()
    wb_out.remove(wb_out.active)

    for ws_in in wb_in.worksheets:
        uuids_to_remove = UUIDS_TO_REMOVE.get(ws_in.title, [])
        move_spec = UUIDS_TO_MOVE.get(ws_in.title)
        ws_out = wb_out.create_sheet(ws_in.title)
        removed, moved = process_sheet(ws_in, ws_out, uuids_to_remove, move_spec)
        print(f"  {ws_in.title}: removed {removed}, moved {moved} rows")

    wb_out.save(output_path)
    print(f"Written to {output_path}")


if __name__ == '__main__':
    main()
