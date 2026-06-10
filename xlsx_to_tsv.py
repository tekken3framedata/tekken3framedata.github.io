#!/usr/bin/env python3
"""Combine all *_framedata.xlsx files in the current directory into a single framedata.tsv."""

import glob
import os

from openpyxl import load_workbook

OUTPUT = 'framedata.tsv'


def main():
    files = sorted(f for f in glob.glob('*_framedata.xlsx') if not os.path.basename(f).startswith('~$'))
    if not files:
        print("No *_framedata.xlsx files found in current directory.")
        return

    header_written = False
    with open(OUTPUT, 'w', encoding='utf-8') as f:
        for xlsx_path in files:
            wb = load_workbook(xlsx_path, read_only=True)
            ws = wb.active
            for row in ws.iter_rows(values_only=True):
                uuid_cell = str(row[0]) if row[0] else ''
                if uuid_cell == 'UUID':
                    if not header_written:
                        cells = [str(cell) if cell is not None else '' for cell in row]
                        f.write('\t'.join(cells) + '\n')
                        header_written = True
                    continue
                if len(uuid_cell) != 36 or uuid_cell.count('-') != 4:
                    continue
                cells = [str(cell) if cell is not None else '' for cell in row]
                f.write('\t'.join(cells) + '\n')
            wb.close()
            print(f"  + {xlsx_path}")

    print(f"-> {OUTPUT}")


if __name__ == '__main__':
    main()
