#!/usr/bin/env python3
"""
Step 5b: Fill missing x markers in the Speed column.

Reads sources/<character>_step5.xlsx. Checks Block Adv, Hit Adv, and
Counter Hit Adv columns for x markers that indicate multi-hit moves.
If those columns have more elements (space-separated) than the Speed column,
the Speed column is padded with trailing x markers to match.

Example: Speed="15", Block Adv="x -2" → Speed becomes "15 x"

Output: sources/<character>_step5b.xlsx

Usage:
    python3 step_5b_hit_markers.py
"""

import glob
import os

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font


ADV_COLUMNS = ('Block Adv', 'Hit Adv', 'Counter Hit Adv')


def fix_hit_markers(rows):
    """Pad Speed column with x markers to match adv column element count."""
    fixed = 0
    for row in rows:
        hit = str(row.get('Speed', '') or '')
        if not hit:
            continue

        hit_parts = hit.split()
        expected_count = 0

        for col in ADV_COLUMNS:
            val = str(row.get(col, '') or '')
            if val:
                parts = val.split()
                if any(p == 'x' for p in parts):
                    expected_count = max(expected_count, len(parts))

        if expected_count > len(hit_parts):
            padding = ['x'] * (expected_count - len(hit_parts))
            row['Speed'] = ' '.join(hit_parts + padding)
            fixed += 1

    return fixed


def parse_sections(ws):
    """Parse a merged worksheet into sections."""
    sections = []
    row = 1
    while row <= ws.max_row:
        cell = ws.cell(row=row, column=1)
        if cell.value and cell.font.bold:
            next_cell = ws.cell(row=row + 1, column=1) if row + 1 <= ws.max_row else None
            if next_cell and next_cell.value == 'Command' and next_cell.font.bold:
                heading = cell.value
                row += 1

                header_cells = []
                for col in range(1, ws.max_column + 1):
                    val = ws.cell(row=row, column=col).value
                    if val:
                        header_cells.append(val)
                row += 1

                data_rows = []
                while row <= ws.max_row:
                    first_cell = ws.cell(row=row, column=1)
                    if first_cell.value is None and all(
                        ws.cell(row=row, column=c).value is None
                        for c in range(1, len(header_cells) + 1)
                    ):
                        row += 1
                        break
                    if first_cell.font.bold:
                        break

                    row_dict = {}
                    for col_idx, col_name in enumerate(header_cells):
                        val = ws.cell(row=row, column=col_idx + 1).value
                        row_dict[col_name] = val or ''
                    data_rows.append(row_dict)
                    row += 1

                sections.append((heading, header_cells, data_rows))
            else:
                row += 1
        else:
            row += 1

    return sections


def autofit_columns(ws):
    """Auto-fit all column widths based on cell content."""
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max_len + 2


def find_characters():
    """Find all characters that have step5 xlsx files."""
    files = glob.glob('sources/*_step5.xlsx')
    return sorted(
        os.path.basename(f).replace('_step5.xlsx', '')
        for f in files
        if not os.path.basename(f).startswith('~$')
    )


def main():
    characters = find_characters()
    if not characters:
        print("No step5 files found in sources/")
        return
    print(f"Found characters: {characters}")

    for char in characters:
        input_path = f"sources/{char}_step5.xlsx"
        output_path = f"sources/{char}_step5b.xlsx"

        print(f"\n=== {char.capitalize()} ===")
        wb_in = load_workbook(input_path)
        ws_in = wb_in.active

        sections = parse_sections(ws_in)

        total_fixed = 0
        for _, _, rows in sections:
            total_fixed += fix_hit_markers(rows)

        print(f"  {total_fixed} Speed values padded")

        wb_out = Workbook()
        ws = wb_out.active
        ws.title = "Merged"
        row_num = 1

        for section_name, header_cells, rows in sections:
            cell = ws.cell(row=row_num, column=1, value=section_name)
            cell.font = Font(bold=True)
            row_num += 1

            for col, h in enumerate(header_cells, 1):
                cell = ws.cell(row=row_num, column=col, value=h)
                cell.font = Font(bold=True)
            row_num += 1

            for row_dict in rows:
                for col, col_name in enumerate(header_cells, 1):
                    val = row_dict.get(col_name, '')
                    if val:
                        ws.cell(row=row_num, column=col, value=val)
                row_num += 1

            row_num += 1

        autofit_columns(ws)
        wb_out.save(output_path)
        print(f"  Written to {output_path}")

    print(f"\nDone. Processed {len(characters)} characters.")


if __name__ == '__main__':
    main()
